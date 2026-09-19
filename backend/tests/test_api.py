"""Offline harness: D1 is backed by a real in-memory sqlite3 so the SQL is
genuinely validated; PoYo is faked."""
import io, json, os, pathlib, sqlite3, sys
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1]))
os.environ.update(CF_ACCOUNT_ID="a", CF_D1_DATABASE_ID="d", CF_API_TOKEN="t", POYO_API_KEY="k",
                  PEXELS_API_KEY="p", BLOB_READ_WRITE_TOKEN="b")

DB = sqlite3.connect(":memory:", check_same_thread=False)
DB.row_factory = sqlite3.Row

CHAT_REPLY = {"value": "hello"}
RAISE_TIMEOUT = {"value": False}

class Resp:
    def __init__(self, body, code=200, headers=None):
        self._b, self.status_code = body, code
        if isinstance(body, bytes):
            self.content, self.text = body, ""
        elif isinstance(body, str):
            self.content, self.text = body.encode(), body
        else:
            self.text = json.dumps(body)
            self.content = self.text.encode()
        self.headers = headers or {"content-type": "application/json"}
    # Mirrors requests.Response: a redirect is a redirect status WITH a
    # Location to follow. _safe_get reads these to walk the chain itself
    # (see its docstring), so a fake that lacks them isn't testing the
    # same code path the real client takes.
    @property
    def is_redirect(self):
        return "location" in {k.lower() for k in self.headers} and self.status_code in (301, 302, 303, 307, 308)
    @property
    def is_permanent_redirect(self):
        return "location" in {k.lower() for k in self.headers} and self.status_code in (301, 308)
    def json(self): return self._b
    def close(self): pass
    def iter_content(self, chunk_size=65536):
        for i in range(0, len(self.content), chunk_size):
            yield self.content[i:i + chunk_size]

def fake_post(url, headers=None, json=None, timeout=None, **kw):
    body = json or {}
    if "/d1/database/" in url:
        sql, params = body["sql"], body.get("params", [])
        cur = DB.cursor()
        try:
            cur.execute(sql, params)
        except sqlite3.Error as e:
            return Resp({"success": False, "errors": [str(e)]}, 200)
        rows = [dict(r) for r in cur.fetchall()] if cur.description else []
        DB.commit()
        return Resp({"success": True, "result": [{"results": rows, "meta": {}}]})
    if "/v1/chat/completions" in url or "/v1/responses" in url:
        if RAISE_TIMEOUT["value"]:
            import requests as _requests
            raise _requests.exceptions.ReadTimeout("Read timed out. (read timeout=110)")
        SENT_MESSAGES.append(body.get("messages") or body.get("input") or [])
        txt = CHAT_REPLY["value"]
        if "/v1/responses" in url:
            return Resp({"data": {"output": [{"content": [{"type": "output_text", "text": txt}]}]}})
        return Resp({"data": {"choices": [{"message": {"content": txt}}]}})
    if "/api/generate/submit" in url:
        SUBMITTED.append(body)
        return Resp({"data": {"task_id": "task-123", "status": "running"}})
    raise AssertionError("unexpected POST " + url)

def fake_get(url, headers=None, timeout=None, params=None, **kw):
    if "/api/generate/status/" in url:
        return Resp({"data": {"status": "finished", "progress": 100,
                              "files": [{"file_url": "https://cdn/x.png"}]}})
    if "/v1/models" in url:
        return Resp({"data": []})
    if "api.pexels.com/v1/search" in url:
        return Resp({"total_results": 1, "photos": [{
            "id": 111, "width": 1000, "height": 1000, "url": "https://pexels.com/photo/111",
            "photographer": "Ada Lovelace", "photographer_url": "https://pexels.com/@ada",
            "src": {"large2x": "https://images.pexels.com/111-large2x.jpg", "medium": "https://images.pexels.com/111-medium.jpg"},
        }]})
    if "api.pexels.com/videos/search" in url:
        return Resp({"total_results": 1, "videos": [{
            "id": 222, "width": 1920, "height": 1080, "image": "https://images.pexels.com/222.jpg",
            "url": "https://pexels.com/video/222", "user": {"name": "Grace Hopper", "url": "https://pexels.com/@grace"},
            "video_files": [
                {"file_type": "video/mp4", "width": 640, "link": "https://videos.pexels.com/222-sd.mp4"},
                {"file_type": "video/mp4", "width": 1920, "link": "https://videos.pexels.com/222-hd.mp4"},
            ],
        }]})
    if url in REDIRECTS:
        return Resp(b"", 302, {"location": REDIRECTS[url]})
    if url in BLOBS:
        r = Resp(BLOBS[url])
        if url in BLOB_CONTENT_TYPES:
            r.headers = {"content-type": BLOB_CONTENT_TYPES[url]}
        return r
    if url == FAKE_WEBSITE_URL:
        return Resp(FAKE_WEBSITE_HTML)
    raise AssertionError("unexpected GET " + url)

FAKE_WEBSITE_URL = "https://example-brand.test/"
FAKE_WEBSITE_HTML = """<html><head>
<title>Acme Coffee Co.</title>
<meta name="description" content="Small-batch coffee, roasted for people who read the label.">
<meta property="og:image" content="/static/logo.png">
<style>.hero{color:#1a2b3c;background:#f5e6c8;} .cta{font-family: 'Poppins', sans-serif;}</style>
</head><body>
<p>We roast in small batches every Tuesday. No jargon, no nonsense — just really good coffee
made by people who care about where it comes from and who it reaches.</p>
</body></html>"""

BLOBS = {}
BLOB_CONTENT_TYPES = {}  # url -> content-type, for tests that care what proxy-image sees
REDIRECTS = {}  # url -> Location it 302s to, for the SSRF redirect-chain tests
def fake_put(url, headers=None, data=None, timeout=None, **kw):
    assert "blob.vercel-storage.com" in url
    pathname = url.split("blob.vercel-storage.com/")[-1] + "-rand"
    blob_url = f"https://blob.example/{pathname}"
    BLOBS[blob_url] = data
    return Resp({"url": blob_url, "pathname": pathname, "contentType": headers.get("x-content-type")})

DELETED_BLOBS = []
def fake_delete(url, headers=None, json=None, timeout=None, **kw):
    assert url.endswith("/delete")
    DELETED_BLOBS.extend(json.get("urls", []))
    return Resp({})

SUBMITTED = []
SENT_MESSAGES = []  # every chat payload, so prompt assembly is assertable
import requests
requests.post, requests.get, requests.put, requests.delete = fake_post, fake_get, fake_put, fake_delete

import server
from fastapi.testclient import TestClient
c = TestClient(server.app)

def check(name, cond, extra=""):
    print(("PASS  " if cond else "FAIL  ") + name + ("" if cond else f"  -> {extra}"))
    if not cond: FAILS.append(name)
FAILS = []

# --- schema self-provisioning ---
r = c.get("/api/stats")
check("stats boots & provisions schema", r.status_code == 200, r.text)
tables = {r[0] for r in DB.execute("SELECT name FROM sqlite_master WHERE type='table'")}
check("all tables created", {"posts", "generations", "connections", "brand_kits"} <= tables, tables)
cols = {r[1] for r in DB.execute("PRAGMA table_info(posts)")}
check("posts has assets/format/hashtags", {"assets", "format", "hashtags"} <= cols, cols)

# --- brand kits (multiple, named, one default) ---
# Acme ends this block back as the default kit — later tests (brand hashtag
# merge, palette injection) expect it active, same as when this was a
# singleton.
r = c.get("/api/brand-kits")
check("brand kits default list", r.status_code == 200 and len(r.json()) == 1
      and r.json()[0]["colors"]["dark"]["accent"] == "#E2FF3D" and r.json()[0]["color_mode"] == "dark", r.text)
check("the default kit carries a guideline with sane defaults, not a missing key",
      r.json()[0].get("guideline", {}).get("version") == "v1.0", r.json()[0].get("guideline"))

r = c.post("/api/brand-kits", json={"name": "Acme", "voice": "dry and technical",
                                     "hashtags": ["#acme"], "banned_words": ["synergy"]})
check("brand kit creates and becomes default (first kit)", r.status_code == 200 and r.json()["name"] == "Acme" and r.json()["is_default"] is True, r.text)
acme_id = r.json()["id"]
r = c.get(f"/api/brand-kits/{acme_id}")
check("brand kit persists", r.json()["voice"] == "dry and technical", r.text)
check("brand_prompt renders", "dry and technical" in server.brand_prompt(r.json()), server.brand_prompt(r.json()))

r = c.post("/api/brand-kits", json={"name": "Side project"})
check("second brand kit creates, not default", r.status_code == 200 and r.json()["is_default"] is False, r.text)
side_id = r.json()["id"]
r = c.get("/api/brand-kits")
check("brand kits list has both, default first", len(r.json()) == 2 and r.json()[0]["id"] == acme_id, r.text)

r = c.put(f"/api/brand-kits/{side_id}", json={"is_default": True})
check("set-default switches it", r.status_code == 200 and r.json()["is_default"] is True, r.text)
r = c.get(f"/api/brand-kits/{acme_id}")
check("old default is demoted", r.json()["is_default"] is False, r.text)

r = c.put(f"/api/brand-kits/{acme_id}", json={"is_default": True, "color_mode": "light", "colors": {"light": {"bg": "#FFFFFF"}}})
check("color_mode switches, light palette merges without clobbering dark, and default moves back",
      r.json()["color_mode"] == "light" and r.json()["colors"]["light"]["bg"] == "#FFFFFF"
      and r.json()["colors"]["dark"]["accent"] == "#E2FF3D" and r.json()["is_default"] is True, r.text)

# --- brand guideline: the extra one-page-summary fields, edited and saved
# through the same PUT the rest of the kit already uses ---
r = c.put(f"/api/brand-kits/{acme_id}", json={"guideline": {
    "logo_clear_space": "Keep a margin >= the icon mark's height",
    "logo_dos": ["Use approved source files"], "logo_donts": ["Recolor or stretch"],
    "voice_attributes": ["Confident", "Warm"], "voice_do": "Ship it.", "doc_owner": "Jamie, Brand",
    "version": "v2.0",
}})
check("guideline fields save", r.status_code == 200, r.text)
guideline = r.json()["guideline"]
check("the edited guideline fields persist", guideline["logo_clear_space"].startswith("Keep a margin"), guideline)
check("list fields persist", guideline["logo_dos"] == ["Use approved source files"] and guideline["voice_attributes"] == ["Confident", "Warm"], guideline)
check("fields not touched by this update keep their defaults, not go missing",
      guideline["icon_style"] == "" and guideline["imagery_mood"] == "", guideline)
r = c.get(f"/api/brand-kits/{acme_id}")
check("guideline survives a reload, not just the immediate response", r.json()["guideline"]["doc_owner"] == "Jamie, Brand", r.json()["guideline"])
check("...and other kit fields are untouched by a guideline-only update", r.json()["voice"] == "dry and technical", r.text)

# A second, unrelated partial guideline update must not wipe out what the
# first one set — a real merge onto the saved guideline, not a replacement.
r = c.put(f"/api/brand-kits/{acme_id}", json={"guideline": {"icon_style": "Outline only, 24px grid"}})
guideline2 = r.json()["guideline"]
check("a later partial guideline update doesn't erase an earlier one's fields",
      guideline2["doc_owner"] == "Jamie, Brand" and guideline2["voice_attributes"] == ["Confident", "Warm"], guideline2)
check("...while still applying its own new field", guideline2["icon_style"] == "Outline only, 24px grid", guideline2)

# --- three logo lockups (color, black-on-white, white-on-black) ---
check("a fresh guideline's logos default to three empty, present keys, not a missing dict",
      set(guideline2.get("logos", {}).keys()) == {"color", "black_on_white", "white_on_black"}, guideline2.get("logos"))
r = c.put(f"/api/brand-kits/{acme_id}", json={"guideline": {"logos": {"color": "https://x/color.png"}}})
logos1 = r.json()["guideline"]["logos"]
check("setting one logo variant doesn't blank the other two",
      logos1 == {"color": "https://x/color.png", "black_on_white": "", "white_on_black": ""}, logos1)
r = c.put(f"/api/brand-kits/{acme_id}", json={"guideline": {"logos": {"black_on_white": "https://x/black.png", "white_on_black": "https://x/white.png"}}})
logos2 = r.json()["guideline"]["logos"]
check("a later update to the other two variants keeps the first one already set",
      logos2 == {"color": "https://x/color.png", "black_on_white": "https://x/black.png", "white_on_black": "https://x/white.png"}, logos2)
r = c.get(f"/api/brand-kits/{acme_id}")
check("all three logo lockups survive a reload", r.json()["guideline"]["logos"] == logos2, r.json()["guideline"])

r = c.get("/api/brand-kits/does-not-exist")
check("unknown brand kit 404s", r.status_code == 404, r.text)

r = c.delete(f"/api/brand-kits/{side_id}")
check("delete a non-default kit", r.status_code == 200, r.text)
r = c.get("/api/brand-kits")
check("acme remains, still the default", len(r.json()) == 1 and r.json()[0]["id"] == acme_id and r.json()[0]["is_default"] is True, r.text)

# --- auto-saved text generations ---
CHAT_REPLY["value"] = "1. First idea\n2. Second idea"
r = c.post("/api/ai/ideate", json={"topic": "distribution", "count": 2})
check("ideate returns ideas", r.json()["ideas"] == ["First idea", "Second idea"], r.text)
check("ideate auto-saved", bool(r.json().get("generation_id")), r.text)

CHAT_REPLY["value"] = "A great post about things."
r = c.post("/api/ai/write", json={"brief": "b", "platform": "linkedin"})
check("write auto-saved", bool(r.json().get("generation_id")), r.text)

CHAT_REPLY["value"] = '{"score": 80, "strengths": ["a"], "improvements": ["b"], "hook_rewrite": "x"}'
r = c.post("/api/ai/coach", json={"content": "draft"})
check("coach auto-saved", bool(r.json().get("generation_id")), r.text)

CHAT_REPLY["value"] = '{"title":"T","slides":[{"heading":"h","body":"b"}]}'
r = c.post("/api/ai/visual", json={"template": "carousel", "topic": "x", "count": 3})
check("visual auto-saved", bool(r.json().get("generation_id")), r.text)

r = c.get("/api/generations?group=text")
gens = r.json()
check("history lists text generations", len(gens) >= 4, len(gens))
check("history carries output", any(g["kind"] == "ideate" and "First idea" in (g["output"] or "") for g in gens), gens[:1])

gid = gens[0]["id"]
r = c.put(f"/api/generations/{gid}", json={"title": "Renamed", "favorite": True})
check("generation edit", r.json()["title"] == "Renamed" and r.json()["favorite"] is True, r.text)
check("favorite filter", len(c.get("/api/generations?favorite=true").json()) == 1, c.get("/api/generations?favorite=true").text)
check("generation delete", c.delete(f"/api/generations/{gid}").status_code == 200)
check("deleted is gone", c.get(f"/api/generations/{gid}").status_code == 404)

# --- media generation + reconcile ---
r = c.post("/api/ai/generate", json={"kind": "image", "prompt": "a cat", "options": {"model": "gpt-image-2", "use_brand": True}})
check("image submit", r.json()["task_id"] == "task-123", r.text)
check("brand palette injected into image prompt", "Colour palette" in SUBMITTED[-1]["input"]["prompt"], SUBMITTED[-1])
check("use_brand not leaked to PoYo", "use_brand" not in SUBMITTED[-1]["input"], SUBMITTED[-1])

# The current Studio image/video picker (2026-09-11) — each request forwards
# whatever the frontend sends for that model's own fields; this checks the
# two model-specific bits the backend itself adds or gates: `quality` for
# any gpt-image* id, and generic pass-through of `resolution`/`image_urls`.
r = c.post("/api/ai/generate", json={"kind": "image", "prompt": "a poster",
                                     "options": {"model": "gpt-image-2.5-sunburst", "size": "21:9", "quality": "high", "resolution": "2K"}})
check("gpt-image-2.5-sunburst submit succeeds", r.status_code == 200, r.text)
check("quality forwarded for a gpt-image* model", SUBMITTED[-1]["input"]["quality"] == "high", SUBMITTED[-1])
check("resolution forwarded generically", SUBMITTED[-1]["input"]["resolution"] == "2K", SUBMITTED[-1])
check("size forwarded", SUBMITTED[-1]["input"]["size"] == "21:9", SUBMITTED[-1])
check("submitted against the real model id", SUBMITTED[-1]["model"] == "gpt-image-2.5-sunburst", SUBMITTED[-1])

r = c.post("/api/ai/generate", json={"kind": "image", "prompt": "a scene",
                                     "options": {"model": "seedream-5.0-pro", "size": "16:9", "resolution": "1K"}})
check("seedream-5.0-pro submit succeeds", r.status_code == 200, r.text)
check("a non-gpt-image model gets no quality field (it doesn't declare one)",
      "quality" not in SUBMITTED[-1]["input"], SUBMITTED[-1])
check("its resolution still forwards generically", SUBMITTED[-1]["input"]["resolution"] == "1K", SUBMITTED[-1])

r = c.post("/api/ai/generate", json={"kind": "video", "prompt": "a scene",
                                     "options": {"model": "sora-2-pro-official", "duration": 8,
                                                 "resolution": "1024p", "aspect_ratio": "16:9"}})
check("sora-2-pro-official submit succeeds", r.status_code == 200, r.text)
check("video options forward generically (resolution the base sora model never had)",
      SUBMITTED[-1]["input"]["resolution"] == "1024p" and SUBMITTED[-1]["input"]["duration"] == 8, SUBMITTED[-1])
row = DB.execute("SELECT status FROM generations WHERE task_id='task-123'").fetchone()
check("media row stored running", row["status"] == "running", dict(row))
r = c.get("/api/generations?group=media")
check("reconcile flips to finished", r.json()[0]["status"] == "finished", r.text[:200])
check("reconcile stored files", r.json()[0]["files"][0]["file_url"] == "https://cdn/x.png", r.text[:200])

# --- platform specs + build-post ---
r = c.get("/api/platform-specs")
check("platform specs", "instagram" in r.json()["platforms"], r.text[:120])

CHAT_REPLY["value"] = json.dumps({
    "format": "carousel", "title": "Ship weekly", "hook": "Consistency wins",
    "caption": "Here is why shipping weekly beats going viral.", "hashtags": ["creators", "#growth"],
    "cta": "Follow for more", "alt_text": "carousel", "why_it_works": "specific",
    "visual": {"style": "carousel", "theme": "chalkboard", "title": "Ship weekly",
               "cover_image_prompt": "a workshop bench",
               "slides": [{"heading": "H1", "body": "B1", "image_prompt": "p1"},
                          {"heading": "H2", "body": "B2"},
                          {"heading": "H3", "body": "B3"}]}})
r = c.post("/api/ai/build-post", json={"topic": "shipping weekly", "platform": "instagram", "format": "carousel", "slides": 3})
b = r.json()
check("build-post format", b["format"] == "carousel", b)
check("build-post normalises hashtags", b["hashtags"][0] == "#creators", b["hashtags"])
check("build-post merges brand hashtag", "#acme" in b["hashtags"], b["hashtags"])
# Acme is the default brand kit here (see line ~143) — a real saved kit
# wins over whatever theme the model happened to pick, so its own
# colors/logo/fonts actually show up by default instead of only after a
# manual click on the brand swatch in the Composer.
check("build-post defaults to the real brand kit over the model's theme pick", b["theme"] == "brand", b["theme"])

r = c.post("/api/ai/build-post", json={"topic": "shipping weekly", "platform": "instagram", "format": "carousel",
                                        "slides": 3, "use_brand": False})
check("build-post honors the model's own theme when brand is opted out", r.json()["theme"] == "chalkboard", r.json()["theme"])

check("build-post assets = cover + slides", len(b["assets"]) == 4, [a["spec"]["template"] for a in b["assets"]])
check("cover first", b["assets"][0]["spec"]["template"] == "cover", b["assets"][0])
check("slide indices", [a["spec"]["index"] for a in b["assets"]] == [0, 1, 2, 3], b["assets"])
check("assets carry no data URLs", "data:image" not in json.dumps(b["assets"]))
check("build-post auto-saved", bool(b.get("generation_id")), b)

CHAT_REPLY["value"] = json.dumps({
    "format": "reel", "title": "R", "caption": "cap", "hashtags": [],
    "visual": {"style": "video", "theme": "midnight",
               "script": [{"scene": "s1", "on_screen_text": "T1", "voiceover": "V1", "video_prompt": "vp1"},
                          {"scene": "s2", "on_screen_text": "T2", "voiceover": "V2", "video_prompt": "vp2"}]}})
r = c.post("/api/ai/build-post", json={"topic": "x", "platform": "tiktok", "format": "reel", "slides": 3})
b = r.json()
check("reel assets are scenes", [a["type"] for a in b["assets"]] == ["scene", "scene"], b["assets"])
check("scene keeps video prompt", b["assets"][0]["spec"]["video_prompt"] == "vp1", b["assets"][0])

CHAT_REPLY["value"] = "not json at all"
r = c.post("/api/ai/build-post", json={"topic": "x", "platform": "linkedin"})
check("build-post survives non-JSON", r.status_code == 200 and r.json()["caption"] == "not json at all", r.text[:200])

# --- posts with assets ---
r = c.post("/api/posts", json={"title": "P", "content": "c", "platforms": ["instagram"],
                               "format": "carousel", "hashtags": ["#a"],
                               "assets": [{"type": "visual", "spec": {"template": "cover", "title": "T"}}]})
check("post create with assets", r.status_code == 200 and r.json()["assets"][0]["spec"]["title"] == "T", r.text[:200])
pid = r.json()["id"]
r = c.put(f"/api/posts/{pid}", json={"assets": [{"type": "visual", "spec": {"template": "slide"}}], "format": "reel"})
check("post update assets", r.json()["assets"][0]["spec"]["template"] == "slide" and r.json()["format"] == "reel", r.text[:200])
r = c.get(f"/api/posts/{pid}")
check("post round-trips assets", r.json()["assets"][0]["spec"]["template"] == "slide", r.text[:200])

# --- posts: alt_text and per-platform caption overrides ---
r = c.post("/api/posts", json={
    "title": "Multi-platform", "content": "The Instagram-length version, up to 2200 chars.",
    "platforms": ["instagram", "twitter"], "alt_text": "A photo of the studio at dawn.",
    "content_by_platform": {"twitter": "The 280-char version."},
})
mp = r.json()
check("post create stores alt_text", mp["alt_text"] == "A photo of the studio at dawn.", mp)
check("post create stores a per-platform caption override",
      mp["content_by_platform"] == {"twitter": "The 280-char version."}, mp)
check("an un-overridden platform is simply absent from content_by_platform",
      "instagram" not in mp["content_by_platform"], mp)
mp_id = mp["id"]
r = c.get(f"/api/posts/{mp_id}")
check("alt_text and content_by_platform round-trip", r.json()["alt_text"] == "A photo of the studio at dawn."
      and r.json()["content_by_platform"] == {"twitter": "The 280-char version."}, r.json())
r = c.put(f"/api/posts/{mp_id}", json={"alt_text": "Updated description.", "content_by_platform": {}})
check("alt_text can be updated and content_by_platform cleared",
      r.json()["alt_text"] == "Updated description." and r.json()["content_by_platform"] == {}, r.json())

r = c.post("/api/posts", json={"title": "Single platform", "content": "just one", "platforms": ["instagram"]})
check("alt_text defaults to empty string", r.json()["alt_text"] == "", r.json())
check("content_by_platform defaults to an empty object", r.json()["content_by_platform"] == {}, r.json())

# --- legacy: media library still only shows real media ---
r = c.get("/api/media")
check("media excludes text generations", all(m["task_id"] for m in r.json()), r.text[:200])

# --- template styles + restyle ---
r = c.get("/api/template-styles")
keys = [t["key"] for t in r.json()["templates"]]
check("template styles listed", set(keys) == {"hooks", "story", "listicle", "contrarian", "how_to"}, keys)
check("template styles carry display copy", all("label" in t and "desc" in t for t in r.json()["templates"]))

CHAT_REPLY["value"] = "Nobody remembers safe. Ship the ugly version today."
r = c.post("/api/ai/restyle", json={"content": "You should post consistently.", "template": "contrarian", "platform": "linkedin"})
check("restyle rewrites the draft", r.json()["content"] == "Nobody remembers safe. Ship the ugly version today.", r.text[:200])
check("restyle auto-saved", bool(r.json().get("generation_id")), r.text)
gid = r.json()["generation_id"]
saved = c.get(f"/api/generations/{gid}").json()
check("restyle recorded under its own kind", saved["kind"] == "restyle", saved)
check("restyle keeps original prompt for context", saved["prompt"] == "You should post consistently.", saved)
r = c.get("/api/generations?group=text")
check("restyle shows up in text history", any(g["kind"] == "restyle" for g in r.json()))

# --- stock media (Pexels) ---
server.PEXELS_API_KEY = None
r = c.get("/api/stock/search", params={"q": "sunset", "type": "image"})
check("stock search reports not configured", r.status_code == 500 and "not configured" in r.json()["detail"], r.text)

server.PEXELS_API_KEY = "test-key"
r = c.get("/api/stock/search", params={"q": "sunset", "type": "image"})
b = r.json()
check("stock image search", len(b["results"]) == 1 and b["results"][0]["type"] == "image", b)
check("stock image picks large url + credit", b["results"][0]["url"].endswith("large2x.jpg") and b["results"][0]["credit"] == "Ada Lovelace", b["results"][0])

r = c.get("/api/stock/search", params={"q": "ocean", "type": "video", "orientation": "portrait"})
b = r.json()
check("stock video search", len(b["results"]) == 1 and b["results"][0]["type"] == "video", b)
check("stock video prefers the hd file", b["results"][0]["url"].endswith("hd.mp4"), b["results"][0])
check("stock video carries thumbnail + credit", b["results"][0]["thumbnail"] and b["results"][0]["credit"] == "Grace Hopper", b["results"][0])

r = c.get("/api/stock/search", params={"q": "x", "type": "bogus"})
check("stock search rejects a bad type", r.status_code == 400, r.text)

# --- uploads (Vercel Blob) ---
server.BLOB_READ_WRITE_TOKEN = None
r = c.post("/api/upload", files={"file": ("photo.png", b"fakepngbytes", "image/png")})
check("upload reports not configured", r.status_code == 500 and "not configured" in r.json()["detail"], r.text)

server.BLOB_READ_WRITE_TOKEN = "test-token"
r = c.post("/api/upload", files={"file": ("my photo!!.png", b"fakepngbytes", "image/png")})
up = r.json()
check("upload stores in blob and D1", r.status_code == 200 and up["url"].startswith("https://blob.example/"), r.text)
check("upload infers kind from content-type", up["kind"] == "image", up)
check("upload sanitises the filename", up["filename"] == "my-photo-.png", up["filename"])
check("upload sends the real bytes to blob", BLOBS[up["url"]] == b"fakepngbytes")
image_upload_id = up["id"]

r = c.post("/api/upload", files={"file": ("clip.mp4", b"fakevideobytes", "video/mp4")})
video_upload = r.json()
check("upload of a video infers kind", video_upload["kind"] == "video", video_upload)

r = c.post("/api/upload", files={"file": ("empty.png", b"", "image/png")})
check("upload rejects an empty file", r.status_code == 400, r.text)

r = c.get("/api/uploads")
check("uploads list returns newest first", [u["id"] for u in r.json()][:2] == [video_upload["id"], image_upload_id], r.text[:200])
r = c.get("/api/uploads", params={"kind": "video"})
check("uploads list filters by kind", len(r.json()) == 1 and r.json()[0]["kind"] == "video", r.text[:200])

r = c.delete(f"/api/uploads/{image_upload_id}")
check("upload delete succeeds", r.status_code == 200, r.text)
check("upload delete removes the blob too", up["url"] in DELETED_BLOBS, DELETED_BLOBS)
check("deleted upload is gone from the list", all(u["id"] != image_upload_id for u in c.get("/api/uploads").json()))
r = c.delete(f"/api/uploads/{image_upload_id}")
check("deleting twice 404s", r.status_code == 404, r.text)

# --- custom fonts (a licensed face the user uploads themselves) ---
WOFF2 = b"wOF2" + b"\x00" * 60
TTF = b"\x00\x01\x00\x00" + b"\x00" * 60

r = c.post("/api/fonts", files={"file": ("brittany.otf", b"this is really a text file", "font/otf")})
check("font upload rejects bytes that aren't a font", r.status_code == 400, r.text)

r = c.post("/api/fonts", files={"file": ("Brittany-Regular.woff2", WOFF2, "application/octet-stream")})
font = r.json()
check("font upload accepts a woff2 mislabelled by the browser", r.status_code == 200, r.text)
check("font upload sniffs the real type", font["content_type"] == "font/woff2", font)
check("font name falls back to a readable filename", font["name"] == "Brittany Regular", font)
check("font bytes reach blob storage", BLOBS[font["url"]] == WOFF2)

r = c.post("/api/fonts", files={"file": ("x.ttf", TTF, "font/ttf")}, data={"name": "Moon'}time; body{x"})
check("font name is stripped of anything that could escape a CSS rule",
      r.json()["name"] == "Moon time body x", r.json())
moon_id = r.json()["id"]

r = c.post("/api/fonts", files={"file": ("replacement.woff2", WOFF2 + b"v2", "application/octet-stream")},
           data={"name": "Brittany Regular"})
check("re-uploading a family replaces it rather than duplicating", r.json()["id"] == font["id"], r.json())
check("the replaced blob is cleaned up", font["url"] in DELETED_BLOBS, DELETED_BLOBS)

rows = c.get("/api/fonts").json()
check("fonts list holds one row per family", len(rows) == 2, rows)
check("fonts list is sorted by name", [f["name"] for f in rows] == ["Brittany Regular", "Moon time body x"], rows)

r = c.delete(f"/api/fonts/{moon_id}")
check("font delete succeeds", r.status_code == 200, r.text)
check("deleted font is gone from the list", [f["name"] for f in c.get("/api/fonts").json()] == ["Brittany Regular"])
r = c.delete(f"/api/fonts/{moon_id}")
check("deleting a font twice 404s", r.status_code == 404, r.text)

r = c.post("/api/upload", files={"file": ("face.woff2", WOFF2, "font/woff2")})
check("a font through the generic upload route is labelled as one", r.json()["kind"] == "font", r.json())

# --- an uploaded image as a reference for image generation ---
r = c.post("/api/ai/generate", json={
    "kind": "image", "prompt": "restyle this product shot",
    "options": {"model": "gpt-image-2", "image_urls": [video_upload["url"]]},
})
check("reference image forwarded to PoYo", SUBMITTED[-1]["input"]["image_urls"] == [video_upload["url"]], SUBMITTED[-1])

# a video model's own reference fields pass straight through untouched
r = c.post("/api/ai/generate", json={
    "kind": "video", "prompt": "animate this photo",
    "options": {"model": "seedance-2-fast", "resolution": "720p", "duration": 5, "image_urls": [video_upload["url"]]},
})
check("video image_urls (first-frame) forwarded", SUBMITTED[-1]["input"]["image_urls"] == [video_upload["url"]], SUBMITTED[-1])

# --- music mashup from two uploaded tracks ---
r = c.post("/api/ai/generate", json={
    "kind": "music", "prompt": "blend these into something upbeat",
    "options": {"reference_urls": ["https://blob.example/a.mp3", "https://blob.example/b.mp3"], "mv": "V4_5"},
})
check("mashup switches to generate-mashup", r.json()["task_id"] == "task-123" and SUBMITTED[-1]["model"] == "generate-mashup", SUBMITTED[-1])
check("mashup sends exactly the two tracks", SUBMITTED[-1]["input"]["upload_url_list"] == ["https://blob.example/a.mp3", "https://blob.example/b.mp3"])

r = c.post("/api/ai/generate", json={
    "kind": "music", "prompt": "x", "options": {"reference_urls": ["https://blob.example/only-one.mp3"]},
})
check("mashup rejects anything but exactly two tracks", r.status_code == 400, r.text)

# plain music generation (no references) still works unchanged
r = c.post("/api/ai/generate", json={"kind": "music", "prompt": "a calm piano piece", "options": {"mv": "V4_5"}})
check("plain music generation still uses generate-music", SUBMITTED[-1]["model"] == "generate-music", SUBMITTED[-1])

# --- a PoYo timeout must surface as a clean 504, never a raw exception string ---
RAISE_TIMEOUT["value"] = True
r = c.post("/api/ai/build-post", json={"topic": "x", "platform": "instagram"})
RAISE_TIMEOUT["value"] = False
check("PoYo timeout returns 504", r.status_code == 504, r.text)
check("timeout detail is a clean message, not a raw exception repr", "ReadTimeout" not in r.json()["detail"] and "didn't respond" in r.json()["detail"], r.text)

# --- brand knowledge base ---
CHAT_REPLY["value"] = '{"summary": "How we price and why.", "tags": ["pricing", "positioning"]}'
r = c.post("/api/knowledge", json={
    "brand_kit_id": acme_id, "title": "Pricing philosophy", "kind": "philosophy",
    "content": ("We never compete on price. Our pricing reflects the cost of doing careful work.\n\n"
                "Discounting signals that the original number was invented. We would rather lose a deal "
                "than teach a client that our rates are negotiable.\n\n"
                "Every quote includes a written scope so the number is legible, not a mystery.")})
doc = r.json()
check("knowledge doc ingests pasted text", r.status_code == 200 and doc["title"] == "Pricing philosophy", r.text[:200])
check("ingest auto-summarises and tags", doc["summary"] == "How we price and why." and "pricing" in doc["tags"], doc)
check("ingest records real character count", doc["chars"] > 200, doc["chars"])
chunk_rows = DB.execute("SELECT count(*) FROM knowledge_chunks WHERE doc_id = ?", (doc["id"],)).fetchone()[0]
check("ingest chunks the document", chunk_rows >= 1, chunk_rows)

CHAT_REPLY["value"] = '{"summary": "Our origin story.", "tags": ["story", "founding"]}'
r = c.post("/api/knowledge", json={
    "brand_kit_id": acme_id, "title": "Non-negotiables", "kind": "guideline", "pinned": True,
    "content": "Never claim to be the cheapest. Never promise a delivery date we have not scheduled."})
pinned_doc = r.json()
check("a doc can be pinned at ingest", pinned_doc["pinned"] is True, pinned_doc)

r = c.post("/api/knowledge", json={
    "brand_kit_id": acme_id, "title": "Kiln case study", "kind": "case_study",
    "content": "We rebuilt the Kiln pottery studio's booking flow. Their no-show rate fell by a third "
               "after we added a deposit step and a reminder the morning of the class."})
case_doc = r.json()

r = c.get(f"/api/knowledge?brand_kit_id={acme_id}")
listed = r.json()
check("knowledge list returns the kit's docs", len(listed) == 3, len(listed))
check("pinned docs sort first", listed[0]["pinned"] is True, [d["title"] for d in listed])
check("list omits full content but keeps a preview", all(len(d["content"]) <= 280 for d in listed), [len(d["content"]) for d in listed])

# Retrieval: the pinned doc is always in; the topical one is fetched on merit.
r = c.post("/api/knowledge/search", json={"brand_kit_id": acme_id, "query": "should we offer a discount?"})
found = r.json()
check("search returns a prompt block", r.status_code == 200 and found["chars"] > 0, r.text[:200])
check("pinned knowledge is always included", any(u["pinned"] and u["title"] == "Non-negotiables" for u in found["used"]), found["used"])
check("BM25 retrieves the topically relevant doc", any(u["title"] == "Pricing philosophy" for u in found["used"]), found["used"])
check("an unrelated doc is left out", not any(u["title"] == "Kiln case study" for u in found["used"]), found["used"])

r = c.post("/api/knowledge/search", json={"brand_kit_id": acme_id, "query": "pottery studio booking no-shows"})
check("a different topic retrieves a different doc",
      any(u["title"] == "Kiln case study" for u in r.json()["used"]), r.json()["used"])

# Knowledge reaches an actual generation.
CHAT_REPLY["value"] = "A post about pricing."
r = c.post("/api/ai/write", json={"brief": "why we do not discount", "brand_kit_id": acme_id})
check("write succeeds with knowledge grounding", r.status_code == 200, r.text[:200])
sent = SENT_MESSAGES[-1][0]["content"]
check("the knowledge block reaches the prompt", "BRAND KNOWLEDGE" in sent, sent[-400:])
check("retrieved knowledge text is in the prompt", "compete on price" in sent, sent[-400:])
check("the brand kit still rides along", "BRAND CONTEXT" in sent, sent[-400:])

r = c.post("/api/ai/write", json={"brief": "why we do not discount", "brand_kit_id": acme_id, "use_knowledge": False})
check("use_knowledge=false opts out", "BRAND KNOWLEDGE" not in SENT_MESSAGES[-1][0]["content"])

# A doc scoped to no kit is shared with every kit.
CHAT_REPLY["value"] = '{"summary": "Shared.", "tags": ["shared"]}'
c.post("/api/knowledge", json={"title": "House style", "kind": "guideline", "pinned": True,
                               "content": "Write in British English. Use the Oxford comma."})
r = c.post("/api/knowledge/search", json={"brand_kit_id": side_id, "query": "anything at all"})
check("kit-less docs apply to every kit", any(u["title"] == "House style" for u in r.json()["used"]), r.json()["used"])
r = c.post("/api/knowledge/search", json={"brand_kit_id": side_id, "query": "discount"})
check("another kit's docs stay out", not any(u["title"] == "Pricing philosophy" for u in r.json()["used"]), r.json()["used"])

# Edit, disable, delete.
r = c.put(f"/api/knowledge/{case_doc['id']}", json={"title": "Kiln — booking rebuild", "pinned": True})
check("a doc can be renamed and pinned", r.json()["title"] == "Kiln — booking rebuild" and r.json()["pinned"] is True, r.text[:200])

# A content edit must refresh the stale summary/tags — a pinned doc's
# summary is injected into EVERY generation verbatim, so leaving it alone
# would keep asserting whatever the text said BEFORE the edit.
CHAT_REPLY["value"] = '{"summary": "Original summary about pricing.", "tags": ["pricing"]}'
r = c.post("/api/knowledge", json={"brand_kit_id": acme_id, "title": "Refund policy", "kind": "guideline",
                                   "pinned": True, "content": "We do not offer refunds after 30 days."})
refund_doc = r.json()
check("summary set at creation", refund_doc["summary"] == "Original summary about pricing.", refund_doc)

CHAT_REPLY["value"] = '{"summary": "Updated summary about extended refunds.", "tags": ["refunds", "extended"]}'
r = c.put(f"/api/knowledge/{refund_doc['id']}", json={"content": "We now offer refunds for up to 90 days."})
updated = r.json()
check("a content edit refreshes the summary", updated["summary"] == "Updated summary about extended refunds.", updated)
check("a content edit refreshes the tags", updated["tags"] == ["refunds", "extended"], updated)

# A rename/pin WITHOUT a content change must NOT re-summarise — no wasted
# model call, and the (still accurate) summary is left exactly alone.
sent_before = len(SENT_MESSAGES)
r = c.put(f"/api/knowledge/{refund_doc['id']}", json={"title": "Refund policy v2", "pinned": True})
check("a metadata-only edit doesn't touch the summary", r.json()["summary"] == "Updated summary about extended refunds.", r.json())
check("a metadata-only edit makes no model call", len(SENT_MESSAGES) == sent_before, len(SENT_MESSAGES) - sent_before)

# The refreshed summary is what a generation actually receives, not just
# what the API response shows.
r = c.post("/api/knowledge/search", json={"brand_kit_id": acme_id, "query": "anything — pinned rides along regardless"})
check("the refreshed summary is what generations actually see",
      any(u["title"] == "Refund policy v2" for u in r.json()["used"]), r.json()["used"])
r = c.delete(f"/api/knowledge/{refund_doc['id']}")
check("cleanup: refund doc deleted", r.status_code == 200)
r = c.put(f"/api/knowledge/{case_doc['id']}", json={"enabled": False})
check("a doc can be disabled", r.json()["enabled"] is False, r.text[:200])
r = c.post("/api/knowledge/search", json={"brand_kit_id": acme_id, "query": "pottery studio booking"})
check("a disabled doc is not retrieved", not any(u["title"].startswith("Kiln") for u in r.json()["used"]), r.json()["used"])

r = c.delete(f"/api/knowledge/{case_doc['id']}")
check("a doc can be deleted", r.status_code == 200, r.text)
check("deleting a doc drops its chunks",
      DB.execute("SELECT count(*) FROM knowledge_chunks WHERE doc_id = ?", (case_doc["id"],)).fetchone()[0] == 0)
r = c.get(f"/api/knowledge/{case_doc['id']}")
check("a deleted doc 404s", r.status_code == 404, r.text)

r = c.post("/api/knowledge", json={"title": "Empty", "content": "   "})
check("an empty document is rejected", r.status_code == 400, r.text[:160])

# --- brand kit export: the only way any of it leaves the database ---
r = c.get(f"/api/brand-kits/{acme_id}/export")
check("export succeeds", r.status_code == 200, r.text[:200])
bundle = r.json()
check("export names the kit", bundle["brand_kit"]["id"] == acme_id, bundle["brand_kit"])
check("export includes the kit's own colours and voice",
      bundle["brand_kit"]["voice"] == "dry and technical", bundle["brand_kit"])
doc_titles = {d["title"] for d in bundle["knowledge_docs"]}
check("export includes this kit's own knowledge docs", "Pricing philosophy" in doc_titles, doc_titles)
check("export includes docs shared across every kit (brand_kit_id NULL)", "House style" in doc_titles, doc_titles)
check("export carries full content, not the truncated list-view preview",
      any(d["title"] == "Pricing philosophy" and len(d["content"]) > 280 for d in bundle["knowledge_docs"]),
      [d["chars"] for d in bundle["knowledge_docs"]])
r = c.get("/api/brand-kits/does-not-exist/export")
check("export 404s for an unknown kit", r.status_code == 404, r.text[:160])

# --- knowledge retrieval: precomputed terms + the warm-instance corpus cache ---
# Chunk and doc term counts are stored at ingest and reused unchanged by
# every generation; a real change (content, pin, mute, delete) must still
# invalidate the cache the very next call — never a stale corpus.
CHAT_REPLY["value"] = '{"summary": "How we quote a custom order.", "tags": ["quote", "estimate"]}'
r = c.post("/api/knowledge", json={"brand_kit_id": acme_id, "title": "Quoting process", "kind": "guideline",
                                   "content": "We always quote in writing before starting any custom work."})
quote_doc = r.json()
chunk_row = DB.execute("SELECT terms FROM knowledge_chunks WHERE doc_id = ?", (quote_doc["id"],)).fetchone()
check("a chunk's term counts are precomputed at ingest",
      chunk_row is not None and json.loads(chunk_row["terms"]).get("quote", 0) > 0, chunk_row and chunk_row["terms"])
doc_row = DB.execute("SELECT meta_terms FROM knowledge_docs WHERE id = ?", (quote_doc["id"],)).fetchone()
check("a doc's title/summary/tags term counts are precomputed at ingest",
      json.loads(doc_row["meta_terms"]).get("estimate", 0) > 0, doc_row["meta_terms"])

r = c.post("/api/knowledge/search", json={"brand_kit_id": acme_id, "query": "a heron never landed here"})
check("an unrelated topic does not retrieve the new doc",
      not any(u["title"] == "Quoting process" for u in r.json()["used"]), r.json()["used"])
r = c.post("/api/knowledge/search", json={"brand_kit_id": acme_id, "query": "how should we quote a custom order"})
check("the new doc is retrieved once its own words are searched",
      any(u["title"] == "Quoting process" for u in r.json()["used"]), r.json()["used"])

# A rewrite must be found on its NEW wording immediately — proves the corpus
# cache is invalidated by the edit rather than serving the pre-edit chunk.
r = c.put(f"/api/knowledge/{quote_doc['id']}", json={"content": "We never start a xylophone-widget commission without a deposit."})
check("content edit succeeds", r.status_code == 200, r.text[:200])
r = c.post("/api/knowledge/search", json={"brand_kit_id": acme_id, "query": "xylophone widget deposit"})
check("the edited content is retrieved right after the edit, not the stale version",
      any(u["title"] == "Quoting process" for u in r.json()["used"]), r.json()["used"])
new_chunk = DB.execute("SELECT terms FROM knowledge_chunks WHERE doc_id = ?", (quote_doc["id"],)).fetchone()
new_terms = json.loads(new_chunk["terms"])
check("re-chunking on edit recomputes term counts for the NEW text",
      "xylophone-widget" in new_terms and "quote" not in new_terms, new_terms)

r = c.delete(f"/api/knowledge/{quote_doc['id']}")
check("cleanup: quoting doc deleted", r.status_code == 200)

# --- knowledge ingest: every file type a person actually has ---
# The source_type is always "file"; what it really is gets decided by reading
# the bytes, so a mislabelled content_type must not change the outcome.
CHAT_REPLY["value"] = '{"summary": "A file.", "tags": ["file"]}'


def _ingest_file(name, data, content_type, title=None):
    up = c.post("/api/upload", files={"file": (name, data, content_type)}).json()
    body = {"brand_kit_id": acme_id, "kind": "note", "source_type": "file", "source_url": up["url"]}
    if title:
        body["title"] = title
    return c.post("/api/knowledge", json=body)


md_body = b"# Craft notes\n\nWe **sand** every edge by hand.\n\n- No shortcuts\n- No filler\n"
r = _ingest_file("craft-notes.md", md_body, "text/markdown")
md_doc = r.json()
check("a .md file is ingested", r.status_code == 200, r.text[:200])
check("markdown text is kept intact", "sand" in md_doc["content"] and "No shortcuts" in md_doc["content"], md_doc["content"][:200])
check("a .md file records source_kind=text", md_doc["source_kind"] == "text", md_doc["source_kind"])
check("the title falls back to the filename", md_doc["title"] == "Craft notes", md_doc["title"])

r = _ingest_file("values.txt", "Honesty over hype.\nShip when it's right.".encode(), "text/plain", "Values")
check("a .txt file is ingested", r.status_code == 200 and "Honesty over hype" in r.json()["content"], r.text[:200])

# A browser that can't guess the type sends octet-stream — sniffing must win.
r = _ingest_file("mislabelled.md", b"## Studio rules\n\nAlways name the maker.", "application/octet-stream")
check("a text file mislabelled as octet-stream is still read",
      r.status_code == 200 and "name the maker" in r.json()["content"], r.text[:200])

from docx import Document as _DocxDocument
docx_buf = io.BytesIO()
_doc = _DocxDocument()
_doc.add_heading("Brand guidelines", level=1)
_doc.add_paragraph("Never describe us as cheap.")
_table = _doc.add_table(rows=1, cols=2)
_table.rows[0].cells[0].text = "Primary"
_table.rows[0].cells[1].text = "Ink black"
_doc.save(docx_buf)
r = _ingest_file("guidelines.docx", docx_buf.getvalue(),
                 "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
docx_doc = r.json()
check("a .docx file is ingested", r.status_code == 200, r.text[:200])
check("docx paragraphs are read", "Never describe us as cheap" in docx_doc["content"], docx_doc["content"][:300])
check("docx table cells are read", "Primary | Ink black" in docx_doc["content"], docx_doc["content"][:300])
check("a .docx file records source_kind=docx", docx_doc["source_kind"] == "docx", docx_doc["source_kind"])

r = _ingest_file("about.html", b"<html><head><style>p{color:red}</style></head><body><h1>Our story</h1>"
                               b"<p>Founded in a garage.</p><script>var x=1</script></body></html>", "text/html")
html_doc = r.json()
check("an .html file is ingested", r.status_code == 200, r.text[:200])
check("html tags are stripped", "<p>" not in html_doc["content"] and "Founded in a garage" in html_doc["content"], html_doc["content"][:200])
check("html script and style contents are dropped",
      "var x" not in html_doc["content"] and "color:red" not in html_doc["content"], html_doc["content"][:200])

r = _ingest_file("prices.csv", b"tier,price\nstarter,49\nstudio,149\n", "text/csv")
check("a .csv file is ingested", r.status_code == 200 and "studio,149" in r.json()["content"], r.text[:200])

r = _ingest_file("note.rtf", rb"{\rtf1\ansi\deff0 {\fonttbl{\f0 Helvetica;}}\f0\fs24 Slow work, made once.\par}",
                 "application/rtf")
check("an .rtf file is ingested", r.status_code == 200 and "Slow work, made once" in r.json()["content"], r.text[:200])

# Files we genuinely cannot read say so by name, with the way out.
r = _ingest_file("legacy.doc", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1legacy binary", "application/msword")
check("a legacy .doc is refused with a useful message",
      r.status_code == 400 and ".docx" in r.text, r.text[:200])
r = _ingest_file("photo.png", b"\x89PNG\r\n\x1a\n\x00\x00binarydata", "image/png")
check("an image is refused as knowledge", r.status_code == 400, r.text[:200])

# --- brand kit: logo upload + a style field ---
r = c.post("/api/upload", files={"file": ("logo.png", b"fakepngbytes", "image/png")})
logo_upload = r.json()
r = c.post("/api/brand-kits", json={"name": "Logo test kit"})
logo_kit_id = r.json()["id"]
r = c.put(f"/api/brand-kits/{logo_kit_id}", json={"logo_url": logo_upload["url"], "style": "Minimal, high-contrast, lots of negative space."})
check("brand kit accepts a logo_url and a style field", r.json()["logo_url"] == logo_upload["url"] and "negative space" in r.json()["style"], r.text[:200])
r = c.get(f"/api/brand-kits/{logo_kit_id}")
check("brand kit style persists", "negative space" in r.json()["style"], r.text[:200])
check("brand_prompt includes style", "Visual style" in server.brand_prompt(r.json()))

# --- brand kit analysis: image (real pixel colors, no AI) ---
from PIL import Image as _Image
buf = io.BytesIO()
img = _Image.new("RGB", (40, 20))
for x in range(40):
    for y in range(20):
        img.putpixel((x, y), (10, 10, 10) if x < 20 else (226, 255, 61))  # half dark, half lime
img.save(buf, format="PNG")
r = c.post("/api/upload", files={"file": ("brandimg.png", buf.getvalue(), "image/png")})
img_upload = r.json()

r = c.post("/api/brand-kit/analyze", json={"source_type": "image", "source_url": img_upload["url"]})
b = r.json()
check("image analysis extracts real colors", set(b["colors"].keys()) == {"bg", "fg", "accent", "sub"}, b["colors"])
check("image analysis colors are actually the two pixel colors", set(b["colors"].values()) <= {"#0a0a0a", "#e2ff3d"}, b["colors"])
check("image analysis sets logo_url to the image itself", b["logo_url"] == img_upload["url"], b)
check("image analysis is honest about not reading fonts/voice", "can't be read from a picture" in b["source_note"], b["source_note"])
check("brand analysis auto-saved", bool(b.get("generation_id")), b)

# --- brand kit analysis: SVG (real markup parsing, no AI) ---
svg_bytes = b"""<svg xmlns="http://www.w3.org/2000/svg" width="100" height="100">
<rect width="100" height="100" fill="#1a2b3c"/><circle cx="50" cy="50" r="30" fill="#ffcc00"/>
<text font-family="Poppins, sans-serif" fill="#ffffff">Acme</text></svg>"""
r = c.post("/api/upload", files={"file": ("logo.svg", svg_bytes, "image/svg+xml")})
svg_upload = r.json()
r = c.post("/api/brand-kit/analyze", json={"source_type": "svg", "source_url": svg_upload["url"]})
b = r.json()
check("svg analysis finds real hex colors", "#1a2b3c" in b["colors"].values(), b["colors"])
check("svg analysis finds the real font-family", b["fonts"]["display"] == "Poppins", b["fonts"])
check("svg analysis sets logo_url to the svg itself", b["logo_url"] == svg_upload["url"])

# --- brand kit analysis: PDF (real text extraction feeds a text-only AI call) ---
minimal_pdf = (
    b"%PDF-1.4\n1 0 obj<< /Type /Catalog /Pages 2 0 R >>endobj\n"
    b"2 0 obj<< /Type /Pages /Kids [3 0 R] /Count 1 >>endobj\n"
    b"3 0 obj<< /Type /Page /Parent 2 0 R /MediaBox [0 0 200 200] "
    b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>endobj\n"
    b"4 0 obj<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>endobj\n"
    b"5 0 obj<< /Length 60 >>\nstream\nBT /F1 18 Tf 10 100 Td (Bold modern coffee brand) Tj ET\nendstream\nendobj\n"
    b"xref\n0 6\n0000000000 65535 f \n"
    b"trailer<< /Size 6 /Root 1 0 R >>\nstartxref\n0\n%%EOF"
)
r = c.post("/api/upload", files={"file": ("deck.pdf", minimal_pdf, "application/pdf")})
pdf_upload = r.json()
CHAT_REPLY["value"] = '{"voice": "Direct, technical, confident.", "style": "Bold and modern.", "suggested_name": "Acme"}'
r = c.post("/api/brand-kit/analyze", json={"source_type": "pdf", "source_url": pdf_upload["url"]})
b = r.json()
check("pdf analysis infers voice from real extracted text", b["voice"] == "Direct, technical, confident.", b)
check("pdf analysis infers a suggested name", b["detected_name"] == "Acme", b)
check("pdf source_note mentions page count", "1-page" in b["source_note"] or "1 page" in b["source_note"], b["source_note"])

# --- brand kit analysis: URL (real page fetch, regex colors, text-only AI) ---
CHAT_REPLY["value"] = '{"voice": "Warm, unpretentious, a little wry.", "style": "Earthy and handcrafted.", "suggested_name": "Acme Coffee Co."}'
r = c.post("/api/brand-kit/analyze", json={"source_type": "url", "source_url": FAKE_WEBSITE_URL})
b = r.json()
check("url analysis reads real page colors", "#1a2b3c" in b["colors"].values(), b["colors"])
check("url analysis reads a real font-family from CSS", b["fonts"]["display"] == "Poppins", b["fonts"])
check("url analysis resolves a relative og:image to an absolute logo_url", b["logo_url"] == "https://example-brand.test/static/logo.png", b["logo_url"])
check("url analysis infers voice from the real page text", "wry" in b["voice"], b)

r = c.post("/api/brand-kit/analyze", json={"source_type": "bogus", "source_url": "x"})
check("brand analysis rejects a bad source_type", r.status_code == 400, r.text)

# --- templates from a file: PPTX (real per-slide text, abstracted by AI) ---
from pptx import Presentation as _Presentation
pptx_buf = io.BytesIO()
prs = _Presentation()
layout = prs.slide_layouts[1]
s1 = prs.slides.add_slide(layout)
s1.shapes.title.text = "Why consistency wins"
s1.placeholders[1].text_frame.text = "Most people quit in week 6."
s2 = prs.slides.add_slide(layout)
s2.shapes.title.text = "The 20-week curve"
s2.placeholders[1].text_frame.text = "Growth looks flat until it doesn't."
prs.save(pptx_buf)
r = c.post("/api/upload", files={"file": ("deck.pptx", pptx_buf.getvalue(),
                                          "application/vnd.openxmlformats-officedocument.presentationml.presentation")})
pptx_upload = r.json()

CHAT_REPLY["value"] = '{"slides": [{"heading": "State the core claim", "body": "Name the timeframe when most people quit."}, {"heading": "Show the payoff curve", "body": "Describe what changes once you push past that point."}]}'
r = c.post("/api/templates/from-file", json={"source_type": "pptx", "source_url": pptx_upload["url"], "name": "Consistency deck"})
tpl = r.json()
check("pptx template has the right slide count", len(tpl["slides"]) == 2, tpl["slides"])
check("pptx template slides are abstracted, not the deck's literal words", tpl["slides"][0]["heading"] == "State the core claim", tpl["slides"])
check("pptx template format defaults to carousel", tpl["format"] == "carousel", tpl)
pptx_template_id = tpl["id"]

# --- templates from a file: PDF (falls back to real text if AI abstraction misbehaves) ---
CHAT_REPLY["value"] = "not valid json"
r = c.post("/api/templates/from-file", json={"source_type": "pdf", "source_url": pdf_upload["url"]})
tpl = r.json()
pdf_template_id = tpl["id"]
check("pdf template falls back to the real page text when AI output can't be parsed", tpl["slides"][0]["heading"] == "Page 1", tpl["slides"])
# This fixture PDF has no embedded image (no colors to sample) but does
# declare a real font resource (/BaseFont /Helvetica) — that alone should
# still be enough to build a usable layout, not require both.
check("a PDF with a real font but no image still gets a usable layout",
      bool(tpl.get("layouts")), tpl.get("layouts"))
check("...carrying the PDF's own real font, not a generic default",
      any(e.get("fontFamily") == "Helvetica" for e in tpl["layouts"].get("cover", [])), tpl["layouts"])

# --- templates from a file: PDF vector colors take priority over raster ---
# A Canva/Illustrator/Figma-style export paints its real background and text
# with vector fill-color operators (`rg` + `re f`), not necessarily via an
# embedded raster image — and even when an image IS embedded, it can be an
# unrelated decorative overlay unrepresentative of the page's real design
# (the bug this fixture guards against). This content stream fills a
# full-page rect light blue, then draws text in dark green.
vector_content = (
    b".7725 .8706 1 rg\n0 0 200 200 re f\n"
    b".1961 .2941 .2431 rg\nBT /F1 18 Tf 10 100 Td (Bold modern coffee brand) Tj ET\n"
)
vector_pdf = (
    b"%PDF-1.4\n1 0 obj<< /Type /Catalog /Pages 2 0 R >>endobj\n"
    b"2 0 obj<< /Type /Pages /Kids [3 0 R] /Count 1 >>endobj\n"
    b"3 0 obj<< /Type /Page /Parent 2 0 R /MediaBox [0 0 200 200] "
    b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>endobj\n"
    b"4 0 obj<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>endobj\n"
    b"5 0 obj<< /Length " + str(len(vector_content)).encode() + b" >>\nstream\n" + vector_content + b"endstream\nendobj\n"
    b"xref\n0 6\n0000000000 65535 f \n"
    b"trailer<< /Size 6 /Root 1 0 R >>\nstartxref\n0\n%%EOF"
)
vector_colors = server._pdf_extract(vector_pdf)["vector_colors"]
check("a full-page filled rect (not merely a clip path) is recognized as the background",
      vector_colors["bg"] == "#c5deff", vector_colors)
check("the color active while text is drawn is recognized as the foreground",
      vector_colors["fg"] == "#324b3e", vector_colors)

r = c.post("/api/upload", files={"file": ("vector.pdf", vector_pdf, "application/pdf")})
vector_pdf_upload = r.json()
CHAT_REPLY["value"] = "not valid json"
r = c.post("/api/templates/from-file", json={"source_type": "pdf", "source_url": vector_pdf_upload["url"]})
vector_tpl = r.json()
check("a template built from a PDF with real vector colors uses them, not a generic default",
      vector_tpl["bg_colors"]["cover"] == "#c5deff", vector_tpl.get("bg_colors"))
check("...and the vector-extracted text color reaches the layout too",
      any(e.get("color") == "#324b3e" for e in vector_tpl["layouts"].get("cover", [])), vector_tpl["layouts"])

r = c.post("/api/brand-kit/analyze", json={"source_type": "pdf", "source_url": vector_pdf_upload["url"]})
vector_brand = r.json()
check("brand-kit PDF analysis also prefers real vector colors over raster sampling",
      vector_brand["colors"]["bg"] == "#c5deff", vector_brand["colors"])
check("the source note is honest about where the palette came from",
      "vector" in vector_brand["source_note"].lower(), vector_brand["source_note"])

# A rectangle traced only to set a clip region (`re W n`, never painted)
# must not be mistaken for a painted background.
clip_only_content = (
    b"0 0 200 200 re W n\n"
    b".9 .1 .1 rg\n50 50 20 20 re f\n"
)
clip_only_pdf = (
    b"%PDF-1.4\n1 0 obj<< /Type /Catalog /Pages 2 0 R >>endobj\n"
    b"2 0 obj<< /Type /Pages /Kids [3 0 R] /Count 1 >>endobj\n"
    b"3 0 obj<< /Type /Page /Parent 2 0 R /MediaBox [0 0 200 200] /Contents 5 0 R >>endobj\n"
    b"5 0 obj<< /Length " + str(len(clip_only_content)).encode() + b" >>\nstream\n" + clip_only_content + b"endstream\nendobj\n"
    b"xref\n0 6\n0000000000 65535 f \n"
    b"trailer<< /Size 6 /Root 1 0 R >>\nstartxref\n0\n%%EOF"
)
clip_colors = server._pdf_extract(clip_only_pdf)["vector_colors"]
check("a clip-only rectangle (re W n, never filled) isn't mistaken for a painted background",
      clip_colors["bg"] == "#e61a1a", clip_colors)

# A PDF with no `rg` fill operators at all (a scanned/photo-only page) has
# no vector colors to offer, so callers correctly fall back to raster
# sampling of an embedded image instead.
check("a PDF with no vector fill operators reports no vector colors",
      server._pdf_extract(minimal_pdf)["vector_colors"] is None, minimal_pdf)

# A source image that's uniformly dark (a moody photo, no real bright color
# anywhere) shouldn't produce invisible text — "brightest of six dark
# browns" still needs to fall back to white against a near-black background.
low_contrast = server._suggest_palette(["#010000", "#762a1f", "#532619", "#37110c", "#130805", "#8b3e28"])
check("a uniformly-dark source palette gets a genuinely legible foreground, not just 'the brightest of the dark ones'",
      low_contrast["fg"] == "#ffffff", low_contrast)
# The reverse case: a uniformly light/pastel source (a beige scrapbook
# background) needs dark text, not "the darkest of the light ones".
low_contrast_light = server._suggest_palette(["#fdf6ec", "#f7e7d0", "#f0dcc0", "#ebd2ae", "#e6c89c"])
check("a uniformly-light source palette also gets a genuinely legible foreground",
      server._contrast_ratio(low_contrast_light["fg"], low_contrast_light["bg"]) >= 4.5, low_contrast_light)
# A palette that already has real contrast shouldn't be second-guessed —
# and every extracted palette, regardless of source, ends up legible.
good_contrast = server._suggest_palette(["#0a0a0a", "#f5f5f5"])
check("a palette with real contrast keeps its own foreground unchanged",
      good_contrast["fg"] in ("#0a0a0a", "#f5f5f5"), good_contrast)
for palette in (low_contrast, low_contrast_light, good_contrast):
    check("the final bg/fg pairing is always genuinely legible",
          server._contrast_ratio(palette["fg"], palette["bg"]) >= 4.5, palette)

CHAT_REPLY["value"] = json.dumps({
    "format": "single", "title": "T", "caption": "cap", "hashtags": [],
    "visual": {"style": "photo", "title": "Real headline", "image_prompt": "x"}})
r = c.post("/api/ai/build-post", json={"topic": "x", "platform": "instagram", "format": "single",
                                        "custom_template_id": pdf_template_id, "use_brand": False})
built = r.json()
cover_els = built["assets"][0]["spec"].get("elements") or []
check("a post built from an uploaded PDF template gets the PDF's real font applied",
      any(e.get("fontFamily") == "Helvetica" for e in cover_els), cover_els)

# A plain "single" + "photo" build with no template picked — the ordinary
# path for a one-image Instagram post — used to come back with zero assets:
# the model is asked for its image idea under "cover_image_prompt" (the same
# key the carousel branch reads for its own cover), but _plan_to_assets only
# ever looked for "image_prompt", a key nothing here produces. With no
# visual asset on the canvas there was nothing for "Save as template" to
# save, which is what "not saving as template after building" actually was.
CHAT_REPLY["value"] = json.dumps({
    "format": "single", "title": "T", "caption": "cap", "hashtags": [],
    "visual": {"style": "photo", "title": "A plain single post", "cover_image_prompt": "a warm coffee shop at dawn"}})
r = c.post("/api/ai/build-post", json={"topic": "coffee", "platform": "instagram", "format": "single", "use_brand": False})
built = r.json()
check("a plain single/photo build (no template) produces a visual asset to save",
      len(built["assets"]) == 1, built)
check("...carrying the model's own cover image prompt",
      built["assets"][0]["spec"].get("image_prompt") == "a warm coffee shop at dawn", built["assets"][0]["spec"])

# --- templates from a file: image (palette only, no slides) ---
r = c.post("/api/templates/from-file", json={"source_type": "image", "source_url": img_upload["url"]})
tpl = r.json()
check("image template has no slide structure", tpl["slides"] == [], tpl)
check("image template format is single", tpl["format"] == "single", tpl)
check("image template carries the real extracted palette", set(tpl["colors"].values()) <= {"#0a0a0a", "#e2ff3d"}, tpl["colors"])

r = c.get("/api/templates/custom")
listed = r.json()
saved_only = [t for t in listed if not t.get("builtin")]
starters = [t for t in listed if t.get("builtin")]
check("custom templates list returns saved templates", len(saved_only) == 4, len(saved_only))
check("saved templates come before the starters", [t.get("builtin", False) for t in listed] == [False] * 4 + [True] * len(starters), len(listed))

# --- starter templates ship with the app ---
check("starters are listed alongside saved templates", len(starters) == len(server.STARTER_TEMPLATES) and len(starters) >= 5, len(starters))
check("every starter carries an outline and a description", all(t["slides"] and t["description"] for t in starters))
check("starters cover more than one format", len({t["format"] for t in starters}) >= 3, {t["format"] for t in starters})
check("starter layouts reference colours by role, not baked hexes",
      all("colorRole" in e for t in server.STARTER_TEMPLATES for layout in t["layouts"].values() for e in layout))

# --- template previews: a real thumbnail spec, not just a description ---
for t in starters:
    check(f"starter '{t['name']}' carries a freeform preview", t["preview"] and "elements" in t["preview"], t.get("preview"))
    texts = [e.get("text", "") for e in t["preview"]["elements"] if e["type"] == "text"]
    check(f"starter '{t['name']}' preview uses its own outline as placeholder copy",
          any(txt and txt in (t["slides"][0]["heading"], t["slides"][0]["body"]) for txt in texts), texts)

pptx_tpl = next(t for t in saved_only if t["source_kind"] == "pptx")
check("a converted pptx template carries a preview", pptx_tpl["preview"] is not None, pptx_tpl["preview"])
check("its preview is a real freeform layout extracted from the deck's own design, not a generic placeholder",
      "elements" in pptx_tpl["preview"] and pptx_tpl["preview"].get("bg_color") == "#ffffff", pptx_tpl["preview"])
check("its preview carries the deck's real theme font, not a generic default",
      any(e.get("fontFamily") == "Calibri" for e in pptx_tpl["preview"]["elements"]), pptx_tpl["preview"])
check("its preview shows the template's own (abstracted) outline text, not blank",
      any(e.get("text") == pptx_tpl["slides"][0]["heading"] for e in pptx_tpl["preview"]["elements"]), pptx_tpl["preview"])
check("the template carries a real layout usable at generation time",
      bool(pptx_tpl.get("layouts")) and bool(pptx_tpl.get("bg_colors")), pptx_tpl.get("layouts"))

image_tpl = next(t for t in saved_only if t["source_kind"] == "image")
check("an image-only template (no slide structure) has no preview to show",
      image_tpl["preview"] is None, image_tpl["preview"])

# --- a real design-tool export: plain text boxes, no placeholders, display
# type with tight leading, and a decorative word drawn wider than the slide ---
from pptx.util import Emu as _Emu, Pt as _Pt

design_buf = io.BytesIO()
dprs = _Presentation()
dprs.slide_width, dprs.slide_height = _Emu(10287000), _Emu(12852400)  # 1080x1350, a 4:5 post
dslide = dprs.slides.add_slide(dprs.slide_layouts[6])  # blank: no title placeholder at all


def _textbox(left, top, width, height, text, size_pt, line_pt, align=None):
    box = dslide.shapes.add_textbox(_Emu(left), _Emu(top), _Emu(width), _Emu(height))
    para = box.text_frame.paragraphs[0]
    run = para.add_run()
    run.text = text
    run.font.size = _Pt(size_pt)
    para.line_spacing = _Pt(line_pt)
    if align is not None:
        para.alignment = align
    return box


# Mirrors the reported deck: a tiny date label sits highest, the headline is
# the biggest phrase, and a one-word script watermark is set larger still and
# stretched to nearly twice the slide width.
_textbox(740258, 722000, 3194231, 303337, "08/09/30", 18.18, 25.46)
_textbox(583956, 2000918, 10614867, 6698173, "Healing\nIs\nNot\nLinear", 145.69, 129.66)
_textbox(6126699, 4913622, 3526397, 3439956, "SOME DAYS YOU WILL FEEL LIGHT, GROUNDED AND HOPEFUL.", 18.18, 25.46)
_textbox(-4362637, 5123488, 18701790, 9042109, "Harper", 526.5, 737.1)
dprs.save(design_buf)
design = server._pptx_extract_design(design_buf.getvalue())
dslide_out = design["slides"][0]
els = dslide_out["elements"]
by_text = {e["text"]: e for e in els}

# 810pt-wide slide, 440px-wide reference card: 145.69pt of headline is 79px
# here. Writing the raw point size (what this used to do) draws it at nearly
# twice the intended size, so the card clips the tail off every heading.
check("a point size is restated against the card's own reference width, not copied raw",
      by_text["Healing\nIs\nNot\nLinear"]["fontSize"] == 79, by_text["Healing\nIs\nNot\nLinear"]["fontSize"])
check("...which holds for body copy on the same slide too",
      by_text["SOME DAYS YOU WILL FEEL LIGHT, GROUNDED AND HOPEFUL."]["fontSize"] == 10,
      by_text["SOME DAYS YOU WILL FEEL LIGHT, GROUNDED AND HOPEFUL."]["fontSize"])
# 129.66pt of leading on a 145.69pt font is 0.89 — display type is routinely
# set tighter than single spacing, and assuming 1.2 overflows its own box.
check("the deck's real line spacing is carried through, not assumed",
      by_text["Healing\nIs\nNot\nLinear"]["lineHeight"] == 0.89, by_text["Healing\nIs\nNot\nLinear"]["lineHeight"])
check("a blank-layout deck with no title placeholder still finds its headline",
      dslide_out["heading"] == "Healing\nIs\nNot\nLinear", dslide_out["heading"])
check("...preferring the phrase over a bigger one-word watermark",
      "Harper" not in dslide_out["heading"], dslide_out["heading"])
check("the headline is the element marked to receive generated copy",
      by_text["Healing\nIs\nNot\nLinear"].get("role_hint") == "title", by_text["Healing\nIs\nNot\nLinear"])
check("...and the paragraph is the one marked for body copy, not the date label",
      by_text["SOME DAYS YOU WILL FEEL LIGHT, GROUNDED AND HOPEFUL."].get("role_hint") == "body", els)
check("a word drawn wider than the slide keeps its bleed rather than being squeezed into a wrap",
      by_text["Harper"]["x"] < 0 and by_text["Harper"]["w"] > 100, by_text["Harper"])

design_layout = server._layout_from_elements(els, True)
title_el = next(e for e in design_layout if e.get("role") == "title")
check("the generated-copy slot is the headline's box, at the headline's size",
      title_el["fontSize"] == 79 and title_el["y"] == 15.57, title_el)
check("...pulled back inside the canvas so fresh copy of any length still wraps into view",
      title_el["x"] >= 0 and title_el["x"] + title_el["w"] <= 100, title_el)
check("a decorative element keeps its own text and stays static",
      any(e.get("text") == "Harper" and not e.get("role") for e in design_layout), design_layout)
check("no internal role hint leaks into the saved layout",
      all("role_hint" not in e for e in design_layout), design_layout)

# --- a saved template's role assignment survives being edited and re-saved ---
# Opening a saved template in the Composer materializes its layout into a
# real, literal-text deck for editing (see templateEdit.js's fillLayout,
# which this simulates) — and role_hint is gone by then, already consumed
# and replaced with `role` the first time _layout_from_elements ran, above.
# Saving that materialized deck again used to re-derive the dynamic slots
# from scratch by Y position, which for this exact deck picks the date
# label (y=5.62) over the actual headline (y=15.57) — corrupting the
# headline's role and the date label's literal text on the very first
# edit-and-save of an uploaded template.
materialized = []
for e in design_layout:
    e = dict(e)
    if e.get("role") == "title":
        e["text"] = "Ship weekly"
    materialized.append(e)
resaved_layout = server._layout_from_elements(materialized, True)
resaved_title = next((e for e in resaved_layout if e.get("role") == "title"), None)
check("re-saving a materialized template keeps the headline as the dynamic title slot",
      resaved_title is not None and resaved_title.get("fontSize") == 79, resaved_title)
check("...instead of reassigning it to whichever element now sits highest on the slide",
      next((e for e in resaved_layout if e.get("text") == "08/09/30"), {}).get("role") is None,
      resaved_layout)
check("...and only one element carries the title role, not two",
      sum(1 for e in resaved_layout if e.get("role") == "title") == 1, resaved_layout)

# --- building a post from an uploaded pptx template actually reproduces its design ---
CHAT_REPLY["value"] = json.dumps({
    "format": "carousel", "title": "T", "caption": "cap", "hashtags": [],
    "visual": {"style": "carousel", "theme": "chalkboard", "title": "Cover",
               "slides": [{"heading": "State the core claim", "body": "Name the timeframe when most people quit."},
                          {"heading": "Show the payoff curve", "body": "Describe what changes once you push past that point."}]}})
r = c.post("/api/ai/build-post", json={"topic": "consistency", "platform": "instagram", "format": "carousel",
                                        "custom_template_id": pptx_tpl["id"], "use_brand": False})
built = r.json()
cover_els = built["assets"][0]["spec"].get("elements") or []
check("a post built from an uploaded pptx template gets real freeform elements, not the model's flat fields",
      bool(cover_els), built["assets"][0]["spec"])
check("...positioned exactly where the source deck had them",
      any(e.get("x") == 5.0 and e.get("y") == 4.0 for e in cover_els), cover_els)
check("...in the source deck's real theme font",
      any(e.get("fontFamily") == "Calibri" for e in cover_els), cover_els)
check("...and the source deck's real background color, not a generic theme",
      built["assets"][0]["spec"].get("bg_color") == "#ffffff", built["assets"][0]["spec"])
check("...with the model's fresh copy poured into the title role",
      any(e.get("text") == "Cover" for e in cover_els), cover_els)

r = c.delete(f"/api/templates/custom/{starters[0]['id']}")
check("a starter can't be deleted", r.status_code == 400 and "can't be deleted" in r.json()["detail"], r.text)

CHAT_REPLY["value"] = json.dumps({
    "format": "carousel", "title": "Ship weekly", "caption": "c", "hashtags": [],
    "visual": {"style": "carousel", "title": "Ship weekly",
               "slides": [{"heading": "H1", "body": "B1"}, {"heading": "H2", "body": "B2"}]}})
r = c.post("/api/ai/build-post", json={"topic": "shipping weekly", "platform": "instagram",
                                       "format": "carousel", "custom_template_id": "starter:bold-hook"})
b = r.json()
check("build-post accepts a starter template id", r.status_code == 200, r.text[:200])
check("a starter lays out every generated slide", all(a["spec"].get("elements") for a in b["assets"]), b["assets"][0]["spec"].keys())
cover_els = b["assets"][0]["spec"]["elements"]
check("layout elements get unique ids", len({e["id"] for e in cover_els}) == len(cover_els))
check("the cover's title element received the generated copy",
      any(e.get("role") == "title" and e.get("text") for e in cover_els), cover_els)
check("static layout copy is kept as authored", any(e.get("text") == "SWIPE →" for e in cover_els), cover_els)
check("empty roles are dropped rather than left as blank boxes",
      all((e.get("text") or "").strip() for e in cover_els if e["type"] == "text"), cover_els)

# --- templates from the Composer, and editing a saved template in place ---
def _text_el(id_, text, font, size, weight, color, y):
    return {"id": id_, "type": "text", "text": text, "x": 8, "y": y, "w": 84, "h": 18,
            "fontFamily": font, "fontSize": size, "fontWeight": weight, "color": color, "align": "left"}

composer_slides_v1 = [
    {"template": "cover", "heading": "Cover", "title": "Cover", "body": "", "bg_color": "#111111",
     "elements": [_text_el("e1", "Cover", "Poppins", 32, 800, "#ffffff", 30)]},
    {"template": "slide", "heading": "Point one", "body": "Detail one", "bg_color": "#111111",
     "elements": [_text_el("e2", "Point one", "Poppins", 24, 800, "#ffffff", 20),
                  _text_el("e3", "Detail one", "Poppins", 16, 400, "#cccccc", 48)]},
]
CHAT_REPLY["value"] = "not valid json"  # falls back to the real outline text, not an AI abstraction
r = c.post("/api/templates/from-composer", json={"name": "My editable template", "format": "carousel", "theme": "midnight", "slides": composer_slides_v1})
check("from-composer saves ok", r.status_code == 200, r.text)
tpl = r.json()
composer_template_id = tpl["id"]
check("from-composer saves a real layout", bool(tpl["layouts"].get("cover")), tpl["layouts"])
check("from-composer saves the slide's own background color", tpl["bg_colors"].get("cover") == "#111111", tpl["bg_colors"])
check("from-composer's real font reaches the layout", any(e.get("fontFamily") == "Poppins" for e in tpl["layouts"]["cover"]), tpl["layouts"])

# Editing: a different theme, background, font, and one more slide than before.
composer_slides_v2 = [
    {"template": "cover", "heading": "New cover", "title": "New cover", "body": "", "bg_color": "#ffe600",
     "elements": [_text_el("e1", "New cover", "Lora", 32, 800, "#111111", 30)]},
    {"template": "slide", "heading": "Point one", "body": "Detail one", "bg_color": "#ffe600",
     "elements": [_text_el("e2", "Point one", "Lora", 24, 800, "#111111", 20)]},
    {"template": "slide", "heading": "Point two", "body": "Detail two", "bg_color": "#ffe600",
     "elements": [_text_el("e4", "Point two", "Lora", 24, 800, "#111111", 20)]},
]
r = c.put(f"/api/templates/custom/{composer_template_id}",
          json={"name": "Renamed template", "format": "carousel", "theme": "whiteboard", "slides": composer_slides_v2})
check("template update succeeds", r.status_code == 200, r.text)
updated = r.json()
check("update overwrites the name", updated["name"] == "Renamed template", updated)
check("update overwrites the theme (style)", updated["theme"] == "whiteboard", updated)
check("update overwrites the background color", updated["bg_colors"]["cover"] == "#ffe600", updated["bg_colors"])
check("update overwrites the font", any(e.get("fontFamily") == "Lora" for e in updated["layouts"]["cover"]), updated["layouts"])
check("update changes the slide count", len(updated["slides"]) == 3, updated["slides"])
check("updating a template edits it in place rather than creating a new one", updated["id"] == composer_template_id, updated)

listed_after_update = c.get("/api/templates/custom").json()
check("editing didn't leave a duplicate row behind",
      sum(1 for t in listed_after_update if t["id"] == composer_template_id) == 1, listed_after_update)

r = c.put("/api/templates/custom/starter:bold-hook", json={"format": "carousel", "slides": []})
check("starter templates can't be edited", r.status_code == 400, r.text)
r = c.put("/api/templates/custom/does-not-exist", json={"format": "carousel", "slides": []})
check("updating an unknown template 404s", r.status_code == 404, r.text)

# Saving an edit must never re-run the AI abstraction over text the user is
# already looking at (materializeTemplateSlides fills a reopened deck's text
# straight from the template's own outline, so on update that text IS the
# outline verbatim) — otherwise every save silently reworded it into a fresh,
# non-deterministic paraphrase, which is what "template text keeps reverting"
# actually was.
CHAT_REPLY["value"] = json.dumps({"slides": [
    {"heading": "State the core claim v1", "body": "Name the timeframe v1"},
    {"heading": "Show the payoff curve v1", "body": "Describe the change v1"},
]})
r = c.post("/api/templates/from-composer", json={"name": "No-drift template", "format": "carousel", "theme": "midnight", "slides": composer_slides_v1})
check("no-drift template saves ok", r.status_code == 200, r.text)
no_drift_id = r.json()["id"]
check("create runs the outline through AI abstraction", r.json()["slides"][0]["heading"] == "State the core claim v1", r.json()["slides"])

# Re-save with the exact same text the abstracted outline already holds (as if
# the template was reopened and saved with no edits), while priming a
# different fake LLM reply that would prove the abstraction step ran again.
resaved_slides = [
    {"template": "cover", "heading": "State the core claim v1", "body": "", "bg_color": "#111111",
     "elements": [_text_el("e1", "State the core claim v1", "Poppins", 32, 800, "#ffffff", 30)]},
    {"template": "slide", "heading": "Show the payoff curve v1", "body": "Describe the change v1", "bg_color": "#111111",
     "elements": [_text_el("e2", "Show the payoff curve v1", "Poppins", 24, 800, "#ffffff", 20),
                  _text_el("e3", "Describe the change v1", "Poppins", 16, 400, "#cccccc", 48)]},
]
CHAT_REPLY["value"] = json.dumps({"slides": [
    {"heading": "DRIFTED v2", "body": "DRIFTED v2"},
    {"heading": "DRIFTED v2", "body": "DRIFTED v2"},
]})
r = c.put(f"/api/templates/custom/{no_drift_id}", json={"format": "carousel", "theme": "midnight", "slides": resaved_slides})
check("no-op resave succeeds", r.status_code == 200, r.text)
resaved = r.json()
check("update does NOT re-run AI abstraction: the text is kept literal, not drifted",
      resaved["slides"][0]["heading"] == "State the core claim v1" and resaved["slides"][1]["body"] == "Describe the change v1",
      resaved["slides"])

# A second no-op save cycle must be just as stable.
r = c.put(f"/api/templates/custom/{no_drift_id}", json={"format": "carousel", "theme": "midnight", "slides": resaved_slides})
resaved_again = r.json()
check("...and stays stable across a second save with no changes",
      resaved_again["slides"] == resaved["slides"], resaved_again["slides"])
r = c.delete(f"/api/templates/custom/{no_drift_id}")
check("cleanup: the no-drift template can be deleted", r.status_code == 200, r.text)

# A post built from the edited template reflects the edit, not the original.
CHAT_REPLY["value"] = json.dumps({
    "format": "carousel", "title": "T", "caption": "cap", "hashtags": [],
    "visual": {"style": "carousel", "title": "Cover", "slides": [{"heading": "h1", "body": "b1"}, {"heading": "h2", "body": "b2"}]},
})
r = c.post("/api/ai/build-post", json={"topic": "x", "platform": "instagram", "format": "carousel",
                                        "custom_template_id": composer_template_id, "use_brand": False})
built = r.json()
cover_els = built["assets"][0]["spec"].get("elements") or []
check("a post built from the edited template uses the edited font", any(e.get("fontFamily") == "Lora" for e in cover_els), cover_els)
check("...and the edited background color", built["assets"][0]["spec"].get("bg_color") == "#ffe600", built["assets"][0]["spec"])


# --- a reel scene's clip is saved and reapplied, even with no custom layout ---
# VideoClipEditor renders for every reel scene regardless of whether it's
# ever entered freeform layout edit, so a scene can carry a stock/uploaded
# clip with no `elements` at all — that must not be gated on `elements` the
# way layout/background are, or the clip is silently dropped on save.
reel_slides_v1 = [
    {"template": "slide", "heading": "Week 1", "body": "You publish. Nobody claps.",
     "clip": {"url": "https://cdn.test/week1.mp4", "credit": "Pexels", "opacity": 0.6, "speed": 1},
     "video_url": "https://cdn.test/week1.mp4"},
    {"template": "slide", "heading": "Week 6", "body": "Three people reply."},
]
r = c.post("/api/templates/from-composer", json={"name": "Reel template", "format": "reel", "theme": "midnight", "slides": reel_slides_v1})
check("a reel with an un-laid-out clip saves ok", r.status_code == 200, r.text)
reel_tpl = r.json()
reel_template_id = reel_tpl["id"]
check("the clip reaches the saved template", reel_tpl.get("clips", {}).get("slide", {}).get("url") == "https://cdn.test/week1.mp4", reel_tpl.get("clips"))
check("...with its edit (opacity) intact, not just the bare url", reel_tpl["clips"]["slide"].get("opacity") == 0.6, reel_tpl["clips"])
check("a scene with no clip contributes nothing (the first scene's clip isn't duplicated onto it)",
      "outro" not in reel_tpl.get("clips", {}) or not reel_tpl["clips"].get("outro"), reel_tpl.get("clips"))

CHAT_REPLY["value"] = json.dumps({
    "format": "reel", "title": "T", "caption": "cap", "hashtags": [],
    "visual": {"style": "video", "script": [{"scene": "s1", "on_screen_text": "a", "voiceover": "b", "video_prompt": "p"},
                                             {"scene": "s2", "on_screen_text": "c", "voiceover": "d", "video_prompt": "q"}]},
})
r = c.post("/api/ai/build-post", json={"topic": "x", "platform": "instagram", "format": "reel",
                                        "custom_template_id": reel_template_id, "use_brand": False})
built = r.json()
check("a reel built from the template gets the saved clip on every scene (falls back to the one saved role)",
      all(a["spec"].get("video_url") == "https://cdn.test/week1.mp4" for a in built["assets"]), built["assets"])

r = c.delete(f"/api/templates/custom/{reel_template_id}")
check("cleanup: the reel template can be deleted", r.status_code == 200, r.text)

# --- reusing a design's media, or asking for new ---
# Picking a design used to mean taking ALL of it: its layout AND the exact
# pictures, footage and background colours it was saved with. That's right
# when you're re-running the same look for a new topic and wrong when the
# pictures were about the old topic. Three independent switches now decide,
# and turning one off must never cost the layout that held the media — the
# frame stays exactly where the design put it, just empty.
media_slides = [
    {"template": "cover", "heading": "Cover", "title": "Cover", "body": "", "bg_color": "#223344",
     "clip": {"url": "https://cdn.test/saved-cover.mp4", "credit": "Pexels"},
     "elements": [
         _text_el("m1", "Cover", "Poppins", 32, 800, "#ffffff", 30),
         {"id": "m2", "type": "image", "url": "https://cdn.test/old-topic.jpg", "x": 5, "y": 55, "w": 60, "h": 30},
         {"id": "m3", "type": "image", "role": "logo", "url": "https://cdn.test/brand-logo.png", "x": 80, "y": 4, "w": 14, "h": 8},
     ]},
    {"template": "slide", "heading": "Point one", "body": "Detail one", "bg_color": "#223344",
     "elements": [_text_el("m4", "Point one", "Poppins", 24, 800, "#ffffff", 20)]},
]
CHAT_REPLY["value"] = "not valid json"
r = c.post("/api/templates/from-composer", json={"name": "Media template", "format": "carousel", "theme": "midnight", "slides": media_slides})
check("a design carrying pictures, a logo and a clip saves ok", r.status_code == 200, r.text)
media_tpl = r.json()
media_template_id = media_tpl["id"]
cover_layout = media_tpl["layouts"]["cover"]
check("the design's own picture reaches the saved layout",
      any(e.get("type") == "image" and e.get("url") == "https://cdn.test/old-topic.jpg" for e in cover_layout), cover_layout)
# The layout builder strips `role` from every element so a demoted text slot
# can't answer to a stale role forever — but a logo has to survive that, or
# there's no way to tell the brand mark apart from a stock photo later.
check("a logo element keeps its role through a save, so it can be told apart from a stock photo",
      any(e.get("role") == "logo" and e.get("url") == "https://cdn.test/brand-logo.png" for e in cover_layout), cover_layout)

def _build_with_media(**flags):
    CHAT_REPLY["value"] = json.dumps({
        "format": "carousel", "title": "T", "caption": "cap", "hashtags": [],
        "visual": {"style": "carousel", "title": "Cover", "slides": [{"heading": "h1", "body": "b1"}, {"heading": "h2", "body": "b2"}]},
    })
    body = {"topic": "x", "platform": "instagram", "format": "carousel",
            "custom_template_id": media_template_id, "use_brand": False}
    body.update(flags)
    resp = c.post("/api/ai/build-post", json=body)
    assert resp.status_code == 200, resp.text
    return resp.json()

def _images(spec, role=None):
    return [e for e in (spec.get("elements") or [])
            if e.get("type") == "image" and (role is None or e.get("role") == role)]

# Default: everything rides along, exactly as before these switches existed.
default_cover = _build_with_media()["assets"][0]["spec"]
check("by default the design's picture still comes along (unchanged behaviour)",
      [e.get("url") for e in _images(default_cover) if e.get("role") != "logo"] == ["https://cdn.test/old-topic.jpg"],
      _images(default_cover))
check("by default the design's background colour still comes along",
      default_cover.get("bg_color") == "#223344", default_cover.get("bg_color"))
check("by default the design's saved clip still comes along",
      default_cover.get("video_url") == "https://cdn.test/saved-cover.mp4", default_cover.get("video_url"))

# Images off: the picture goes, its frame stays, and the slot is flagged so
# the Composer knows to fill THIS one (and only this one) with something new.
new_images = _build_with_media(reuse_template_images=False)["assets"][0]["spec"]
content_imgs = [e for e in _images(new_images) if e.get("role") != "logo"]
check("asking for new images empties the design's picture", [e.get("url") for e in content_imgs] == [""], content_imgs)
check("...but keeps the frame exactly where the design put it",
      [(e.get("x"), e.get("y"), e.get("w"), e.get("h")) for e in content_imgs] == [(5, 55, 60, 30)], content_imgs)
check("...and flags the slot so the app knows to fill it rather than leave it blank",
      all(e.get("needs_image") for e in content_imgs), content_imgs)
check("asking for new images never touches the logo — that's the brand, not this post's content",
      [e.get("url") for e in _images(new_images, role="logo")] == ["https://cdn.test/brand-logo.png"], _images(new_images, role="logo"))
check("a slot the design left blank on purpose isn't flagged (nothing was emptied)",
      not any(e.get("needs_image") for e in _images(default_cover)), _images(default_cover))
check("asking for new images leaves the background and footage alone",
      new_images.get("bg_color") == "#223344" and new_images.get("video_url") == "https://cdn.test/saved-cover.mp4", new_images)

# Backgrounds off: the slide falls back to its own theme colour.
new_bg = _build_with_media(reuse_template_backgrounds=False)["assets"][0]["spec"]
check("asking for a new background drops the design's saved colour", new_bg.get("bg_color") != "#223344", new_bg.get("bg_color"))
check("asking for a new background leaves the pictures alone",
      [e.get("url") for e in _images(new_bg) if e.get("role") != "logo"] == ["https://cdn.test/old-topic.jpg"], _images(new_bg))

# Videos off: no clip is copied, so the scene is an empty slot the Composer's
# existing gap-fill then fetches fresh footage for.
new_video = _build_with_media(reuse_template_videos=False)["assets"][0]["spec"]
check("asking for new footage drops the design's saved clip",
      not new_video.get("video_url") and not new_video.get("clip"), new_video)
check("asking for new footage leaves the layout, pictures and background alone",
      new_video.get("bg_color") == "#223344"
      and [e.get("url") for e in _images(new_video) if e.get("role") != "logo"] == ["https://cdn.test/old-topic.jpg"], new_video)

# All three off at once still keeps the thing you actually picked the design for.
all_new = _build_with_media(reuse_template_images=False, reuse_template_videos=False, reuse_template_backgrounds=False)["assets"][0]["spec"]
check("turning everything off still keeps the design's layout — that's what picking a design means",
      any(e.get("type") == "text" and e.get("fontFamily") == "Poppins" for e in (all_new.get("elements") or [])), all_new.get("elements"))
check("...and still keeps the logo", [e.get("url") for e in _images(all_new, role="logo")] == ["https://cdn.test/brand-logo.png"], _images(all_new, role="logo"))

r = c.delete(f"/api/templates/custom/{media_template_id}")
check("cleanup: the media template can be deleted", r.status_code == 200, r.text)

# --- a design's handle, hashtags and contact line are not copy ---
# Reported: "@connected.mothering is being replaced with a heading". Saving a
# slide as a design turns one or two of its text boxes into copy slots,
# refilled on every post built from it. With nothing else to go on that pick
# is by Y position — and in a real social layout the handle is exactly what
# sits at the top, so the handle became the headline slot and every build
# overwrote it.
mark_slides = [
    {"template": "cover", "heading": "The real headline", "title": "The real headline", "body": "", "bg_color": "#101010",
     "elements": [
         _text_el("k1", "@connected.mothering", "Poppins", 13, 500, "#cccccc", 4),
         _text_el("k2", "The real headline", "Poppins", 32, 800, "#ffffff", 30),
         _text_el("k3", "The supporting line underneath it", "Poppins", 15, 400, "#cccccc", 55),
         _text_el("k4", "#momlife #gentleparenting", "Poppins", 11, 400, "#888888", 88),
         _text_el("k5", "hello@connectedmothering.com", "Poppins", 11, 400, "#888888", 93),
     ]},
    {"template": "slide", "heading": "Point one", "body": "Detail one", "bg_color": "#101010",
     "elements": [_text_el("k6", "Point one", "Poppins", 24, 800, "#ffffff", 20)]},
]
CHAT_REPLY["value"] = "not valid json"
r = c.post("/api/templates/from-composer", json={"name": "Handle design", "format": "carousel", "theme": "midnight", "slides": mark_slides})
check("a design with a handle, hashtags and an email saves ok", r.status_code == 200, r.text)
mark_tpl = r.json()
mark_template_id = mark_tpl["id"]
mark_layout = mark_tpl["layouts"]["cover"]

def _el_text(layout, text):
    return next((e for e in layout if e.get("text") == text), None)

def _roles(layout):
    return {e.get("text", "<dynamic>"): e.get("role") for e in layout if e.get("type") == "text"}

check("the handle keeps its literal text instead of becoming a copy slot",
      _el_text(mark_layout, "@connected.mothering") is not None, _roles(mark_layout))
check("...and carries no copy role, so nothing downstream writes to it",
      (_el_text(mark_layout, "@connected.mothering") or {}).get("role") is None, _roles(mark_layout))
check("the hashtag row is left alone too", _el_text(mark_layout, "#momlife #gentleparenting") is not None, _roles(mark_layout))
check("...as is the contact line", _el_text(mark_layout, "hello@connectedmothering.com") is not None, _roles(mark_layout))
# The point isn't only that the handle survives — the real headline has to be
# the slot instead, or the design just has no headline any more.
title_els = [e for e in mark_layout if e.get("role") == "title"]
body_els = [e for e in mark_layout if e.get("role") == "body"]
check("the actual headline becomes the title slot, not the handle above it",
      len(title_els) == 1 and "text" not in title_els[0], title_els)
check("...and the line under it becomes the body slot",
      len(body_els) == 1 and "text" not in body_els[0], body_els)

CHAT_REPLY["value"] = json.dumps({
    "format": "carousel", "title": "T", "caption": "cap", "hashtags": [],
    "visual": {"style": "carousel", "title": "Generated cover headline",
               "slides": [{"heading": "h1", "body": "b1"}, {"heading": "h2", "body": "b2"}]},
})
r = c.post("/api/ai/build-post", json={"topic": "x", "platform": "instagram", "format": "carousel",
                                        "custom_template_id": mark_template_id, "use_brand": False})
built_cover = r.json()["assets"][0]["spec"]
built_text = [e.get("text") for e in (built_cover.get("elements") or []) if e.get("type") == "text"]
check("a post built from that design still says the handle, not a generated heading",
      "@connected.mothering" in built_text, built_text)
check("...and the hashtags and email survive the build too",
      "#momlife #gentleparenting" in built_text and "hello@connectedmothering.com" in built_text, built_text)
check("...while the generated headline lands in the headline's own box",
      "Generated cover headline" in built_text, built_text)

# The explicit tag, for what reading the text can't know. `fixed` wins both
# ways: it pins an ordinary-looking line, and releases a marky-looking one.
tagged_slides = [
    {"template": "cover", "heading": "Cover", "title": "Cover", "body": "", "bg_color": "#101010",
     "elements": [
         {**_text_el("f1", "Book a discovery call", "Poppins", 14, 600, "#cccccc", 6), "fixed": True},
         _text_el("f2", "Cover", "Poppins", 32, 800, "#ffffff", 30),
         _text_el("f3", "Supporting line", "Poppins", 15, 400, "#cccccc", 55),
     ]},
]
CHAT_REPLY["value"] = "not valid json"
r = c.post("/api/templates/from-composer", json={"name": "Tagged design", "format": "carousel", "theme": "midnight", "slides": tagged_slides})
tagged_layout = r.json()["layouts"]["cover"]
tagged_id = r.json()["id"]
check("a line tagged Keep as is is never made a copy slot, even though it reads as copy",
      (_el_text(tagged_layout, "Book a discovery call") or {}).get("role") is None, _roles(tagged_layout))
check("...and the next line down becomes the headline slot instead",
      len([e for e in tagged_layout if e.get("role") == "title"]) == 1, _roles(tagged_layout))
check("...with the tag itself carried into the design, so it still holds on re-save",
      (_el_text(tagged_layout, "Book a discovery call") or {}).get("fixed") is True, tagged_layout)

released_slides = [
    {"template": "cover", "heading": "Cover", "title": "Cover", "body": "", "bg_color": "#101010",
     "elements": [{**_text_el("r1", "@thehandle", "Poppins", 30, 800, "#ffffff", 20), "fixed": False}]},
]
r = c.post("/api/templates/from-composer", json={"name": "Released design", "format": "carousel", "theme": "midnight", "slides": released_slides})
released_layout = r.json()["layouts"]["cover"]
released_id = r.json()["id"]
check("tagging a handle-looking box Refill releases it, so the auto-read can be overridden either way",
      len([e for e in released_layout if e.get("role") == "title"]) == 1, released_layout)

for _id in (mark_template_id, tagged_id, released_id):
    check("cleanup: the design can be deleted", c.delete(f"/api/templates/custom/{_id}").status_code == 200)



r = c.delete(f"/api/templates/custom/{composer_template_id}")
check("cleanup: the editable template can be deleted", r.status_code == 200, r.text)

r = c.delete(f"/api/templates/custom/{pptx_template_id}")
check("custom template delete succeeds", r.status_code == 200, r.text)
check("deleted template is gone from the list", all(t["id"] != pptx_template_id for t in c.get("/api/templates/custom").json()))
r = c.delete(f"/api/templates/custom/{pptx_template_id}")
check("deleting twice 404s", r.status_code == 404)

r = c.post("/api/templates/from-file", json={"source_type": "bogus", "source_url": "x"})
check("templates-from-file rejects a bad source_type", r.status_code == 400, r.text)

# PDF and PPTX reach the knowledge base through the same generic file path.
CHAT_REPLY["value"] = '{"summary": "A file.", "tags": ["file"]}'
r = c.post("/api/knowledge", json={"brand_kit_id": acme_id, "source_type": "file",
                                   "source_url": pdf_upload["url"], "title": "Deck PDF"})
check("a PDF still ingests through source_type=file",
      r.status_code == 200 and r.json()["source_kind"] == "pdf"
      and "coffee brand" in r.json()["content"], r.text[:200])
r = c.post("/api/knowledge", json={"brand_kit_id": acme_id, "source_type": "file",
                                   "source_url": pptx_upload["url"], "title": "Consistency deck"})
check("a PPTX still ingests through source_type=file",
      r.status_code == 200 and r.json()["source_kind"] == "pptx"
      and "consistency wins" in r.json()["content"].lower(), r.text[:200])
# Old clients that still name the type explicitly keep working.
r = c.post("/api/knowledge", json={"brand_kit_id": acme_id, "source_type": "pdf",
                                   "source_url": pdf_upload["url"], "title": "Legacy call"})
check("an explicit source_type=pdf is still accepted", r.status_code == 200, r.text[:200])

# --- build-post follows a saved template's outline ---
remaining = c.get("/api/templates/custom").json()
image_template_id = next(t["id"] for t in remaining if t["source_kind"] == "image")
r = c.post("/api/ai/build-post", json={"topic": "y", "platform": "instagram", "custom_template_id": "does-not-exist"})
check("build-post 404s on an unknown template id", r.status_code == 404, r.text)

CHAT_REPLY["value"] = json.dumps({
    "format": "single", "title": "T", "caption": "c", "hashtags": [], "visual": {"style": "quote"},
})
r = c.post("/api/ai/build-post", json={"topic": "why coffee matters", "platform": "instagram", "custom_template_id": image_template_id})
b = r.json()
check("build-post with an image template uses the template's theme", b["theme"] == "midnight", b)

# --- SSRF guard: any endpoint that fetches a caller-supplied URL runs through this ---
for bad_url, why in [
    ("ftp://example.com/x", "non-http(s) scheme"),
    ("http://127.0.0.1/", "loopback"),
    ("http://localhost/", "loopback by name"),
    ("http://169.254.169.254/latest/meta-data/", "cloud metadata address"),
    ("http://10.0.0.5/", "private range"),
    ("http://192.168.1.1/", "private range"),
    ("http://[::1]/", "IPv6 loopback"),
]:
    try:
        server._assert_public_http_url(bad_url)
        check(f"SSRF guard rejects {why}", False, bad_url)
    except Exception as e:
        check(f"SSRF guard rejects {why}", isinstance(e, server.HTTPException) and e.status_code == 400, str(e))

# A real public address (literal IP, no DNS needed) is not blocked.
try:
    server._assert_public_http_url("https://8.8.8.8/")
    check("SSRF guard allows a public address", True)
except Exception as e:
    check("SSRF guard allows a public address", False, str(e))

# A hostname that can't resolve at all isn't a reachable target either way —
# every fake domain the rest of this suite uses (blob.example, *.test) relies
# on exactly this not being treated as a block.
try:
    server._assert_public_http_url("https://this-does-not-exist.invalid/")
    check("SSRF guard doesn't block an unresolvable host (can't be an SSRF target)", True)
except Exception as e:
    check("SSRF guard doesn't block an unresolvable host (can't be an SSRF target)", False, str(e))

# --- the guard has to survive a redirect, not just the URL handed in ---
# requests follows redirects on its own, so a public host answering 302 with
# a private Location used to be fetched anyway — past a check that rejects
# that same address outright when passed in directly (asserted above).
REDIRECTS["https://redirector.test/to-metadata"] = "http://169.254.169.254/latest/meta-data/"
try:
    server._fetch_with_cap("https://redirector.test/to-metadata", 1 << 20, 5)
    check("SSRF guard rejects a private address reached via redirect", False, "fetch completed")
except Exception as e:
    check("SSRF guard rejects a private address reached via redirect",
          isinstance(e, server.HTTPException) and e.status_code == 400, f"{type(e).__name__}: {e}")

# ...while an ordinary redirect to a public target still resolves normally.
BLOBS["https://cdn.test/real.png"] = b"redirected-png-bytes"
BLOB_CONTENT_TYPES["https://cdn.test/real.png"] = "image/png"
REDIRECTS["https://redirector.test/to-image"] = "https://cdn.test/real.png"
try:
    body, ctype = server._fetch_with_cap("https://redirector.test/to-image", 1 << 20, 5)
    check("a redirect to a public target is still followed",
          body == b"redirected-png-bytes" and ctype == "image/png", (body[:40], ctype))
except Exception as e:
    check("a redirect to a public target is still followed", False, f"{type(e).__name__}: {e}")

# A loop ends as an error rather than spinning until the request times out.
REDIRECTS["https://redirector.test/loop"] = "https://redirector.test/loop"
try:
    server._fetch_with_cap("https://redirector.test/loop", 1 << 20, 5)
    check("a redirect loop is cut off", False, "fetch completed")
except Exception as e:
    check("a redirect loop is cut off",
          isinstance(e, server.HTTPException) and "too many redirects" in str(e.detail), str(e))

# --- proxy-image: same guard, plus a content-type allowlist ---
r = c.get("/api/proxy-image", params={"url": "http://169.254.169.254/"})
check("proxy-image refuses a private/metadata address", r.status_code == 400, r.text[:160])
BLOBS["https://blob.example/not-an-image"] = b"plain text, not media"
BLOB_CONTENT_TYPES["https://blob.example/not-an-image"] = "text/plain"
r = c.get("/api/proxy-image", params={"url": "https://blob.example/not-an-image"})
check("proxy-image refuses a non-image/video content-type", r.status_code == 400, r.text[:160])
BLOBS["https://blob.example/pic.png"] = b"fakepngbytes"
BLOB_CONTENT_TYPES["https://blob.example/pic.png"] = "image/png"
r = c.get("/api/proxy-image", params={"url": "https://blob.example/pic.png"})
check("proxy-image re-serves an actual image", r.status_code == 200 and r.content == b"fakepngbytes", r.status_code)
check("proxy-image sets a permissive CORS header for canvas export",
      r.headers.get("access-control-allow-origin") == "*", dict(r.headers))

# --- uploads/from-url: the Library's "+" downloading a Stock pick ---
r = c.post("/api/uploads/from-url", params={}, json={"url": "http://169.254.169.254/"})
check("uploads/from-url refuses a private/metadata address", r.status_code == 400, r.text[:160])

BLOBS["https://images.pexels.com/city.jpg"] = b"fake-stock-photo-bytes"
BLOB_CONTENT_TYPES["https://images.pexels.com/city.jpg"] = "image/jpeg"
r = c.post("/api/uploads/from-url", json={"url": "https://images.pexels.com/city.jpg", "filename": "Jane Doe - photo-1"})
check("a stock pick downloads and saves ok", r.status_code == 200, r.text)
stock_upload = r.json()
check("the real bytes were fetched and stored, not just the url referenced", BLOBS[stock_upload["url"]] == b"fake-stock-photo-bytes")
check("its kind is inferred from the real content-type", stock_upload["kind"] == "image", stock_upload)
check("the filename is sanitised the same way a direct upload's is", stock_upload["filename"] == "Jane-Doe---photo-1", stock_upload["filename"])

r = c.get("/api/uploads")
check("the downloaded stock item shows up in the uploads list", any(u["id"] == stock_upload["id"] for u in r.json()), r.text[:200])

BLOBS["https://blob.example/not-media.txt"] = b"plain text, not media"
BLOB_CONTENT_TYPES["https://blob.example/not-media.txt"] = "text/plain"
r = c.post("/api/uploads/from-url", json={"url": "https://blob.example/not-media.txt"})
check("a non-media link is rejected rather than silently added", r.status_code == 400, r.text)

r = c.delete(f"/api/uploads/{stock_upload['id']}")
check("cleanup: the downloaded stock item can be deleted like any upload", r.status_code == 200, r.text)

# --- voice previews: made once, not once per session ---
# Hearing what a voice sounds like used to cost a text-to-speech generation
# every time, because the answer only lived in one page's memory: reload and
# all ten presets were unknown again. The sample never changes, so it only
# ever needs making once.
r = c.get("/api/voice-previews")
check("a fresh install has no saved previews yet", r.status_code == 200 and r.json() == [], r.text[:200])

PREVIEW_LINE = "Hi there — this is a quick preview of how I sound."
BLOBS["https://poyo.cdn/voice-abc.mp3"] = b"fake-voice-sample-bytes"
BLOB_CONTENT_TYPES["https://poyo.cdn/voice-abc.mp3"] = "audio/mpeg"
r = c.post("/api/voice-previews", json={"voice": "Rachel", "source_url": "https://poyo.cdn/voice-abc.mp3", "line": PREVIEW_LINE})
check("a generated sample can be banked", r.status_code == 200, r.text)
saved = r.json()
check("...and comes back with a url to play", bool(saved.get("url")), saved)
# The generator hands back a URL on its OWN CDN, which is theirs to expire. A
# row pointing at one of those is a dead preview button some weeks from now,
# so the bytes have to be re-hosted here — the same thing "downloaded from
# Stock" means for a photo.
check("the sample's real bytes are re-hosted, not just the provider's link bookmarked",
      saved["url"] != "https://poyo.cdn/voice-abc.mp3" and BLOBS.get(saved["url"]) == b"fake-voice-sample-bytes",
      saved)

r = c.get("/api/voice-previews")
listed = r.json()
check("the saved sample is listed for the next session to seed from",
      len(listed) == 1 and listed[0]["voice"] == "Rachel" and listed[0]["url"] == saved["url"], listed)

# The whole point: a second ask for the same voice must not spend another
# upload. Priming a DIFFERENT body at the same source url proves the stored
# row was returned rather than the link re-fetched.
BLOBS["https://poyo.cdn/voice-abc.mp3"] = b"a-second-generation-that-should-never-be-used"
r = c.post("/api/voice-previews", json={"voice": "Rachel", "source_url": "https://poyo.cdn/voice-abc.mp3", "line": PREVIEW_LINE})
check("asking again returns the sample already banked instead of making another",
      r.status_code == 200 and r.json()["url"] == saved["url"], r.text[:200])
check("...and doesn't quietly overwrite it with the second one",
      BLOBS[saved["url"]] == b"fake-voice-sample-bytes", BLOBS.get(saved["url"]))

# A sample is only the right answer for the sentence it was read from. Change
# the line and the stored clip is saying something else.
BLOBS["https://poyo.cdn/voice-new-line.mp3"] = b"read-from-the-new-line"
BLOB_CONTENT_TYPES["https://poyo.cdn/voice-new-line.mp3"] = "audio/mpeg"
r = c.post("/api/voice-previews", json={"voice": "Rachel", "source_url": "https://poyo.cdn/voice-new-line.mp3",
                                        "line": "A completely different sample sentence."})
check("changing the sample line re-records rather than serving the old words",
      r.status_code == 200 and BLOBS.get(r.json()["url"]) == b"read-from-the-new-line", r.text[:200])
check("...and replaces the row rather than adding a second one for the same voice",
      len(c.get("/api/voice-previews").json()) == 1, c.get("/api/voice-previews").json())

r = c.post("/api/voice-previews", json={"voice": "", "source_url": "https://poyo.cdn/voice-abc.mp3"})
check("a preview with no voice attached is refused", r.status_code == 400, r.text[:160])

BLOBS["https://poyo.cdn/not-audio.txt"] = b"plain text"
BLOB_CONTENT_TYPES["https://poyo.cdn/not-audio.txt"] = "text/plain"
r = c.post("/api/voice-previews", json={"voice": "Aria", "source_url": "https://poyo.cdn/not-audio.txt"})
check("something that isn't audio is refused rather than saved as a preview", r.status_code == 400, r.text[:160])

r = c.post("/api/voice-previews", json={"voice": "Aria", "source_url": "http://169.254.169.254/"})
check("the same SSRF guard as every other fetch applies here", r.status_code == 400, r.text[:160])

r = c.delete("/api/voice-previews/Rachel")
check("a sample can be thrown away so the next preview re-records it", r.status_code == 200, r.text)
check("...and it's gone from the list", c.get("/api/voice-previews").json() == [], c.get("/api/voice-previews").json())
r = c.delete("/api/voice-previews/Rachel")
check("deleting one that isn't there 404s", r.status_code == 404, r.text[:160])


# --- app access token: off by default, enforced once set, cron route exempt ---
check("no token configured: every route is open", c.get("/api/stats").status_code == 200)
server.APP_ACCESS_TOKEN = "super-secret"
try:
    check("token configured: an unauthenticated request is rejected",
          c.get("/api/stats").status_code == 401)
    check("token configured: the wrong token is rejected",
          c.get("/api/stats", headers={"Authorization": "Bearer nope"}).status_code == 401)
    check("token configured: the right token is accepted",
          c.get("/api/stats", headers={"Authorization": "Bearer super-secret"}).status_code == 200)
    check("the cron endpoint is exempt from the app token (it has its own CRON_SECRET)",
          c.get("/api/cron/publish-due").status_code == 200)
finally:
    server.APP_ACCESS_TOKEN = None  # restore — every test above and after this block assumes no auth

print("\n" + ("ALL PASS" if not FAILS else f"{len(FAILS)} FAILED: {FAILS}"))
sys.exit(1 if FAILS else 0)
